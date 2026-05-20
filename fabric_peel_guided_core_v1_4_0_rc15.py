"""
fabric_peel_guided_core_v1_4_0_rc15.py

Manuscript-guided peel-trace stability analysis backend for soft textile / flexible-laminate
T-peel data. Version 1.4.0-rc15 intentionally prioritizes paper-replica force-trace metrics over
broad fracture-energy analysis.

Core manuscript profile
-----------------------
- ASTM D2724-style T-peel force-displacement trace analysis.
- Baseline correction using first force value; negative corrected forces clipped to zero.
- Displacement sign-corrected when needed while preserving row order.
- Candidate 25.0 mm windows start after a 5.0 mm offset and must remain outside the
  terminal 20% of displacement range.
- Drift gate: |m| L_win / Fbar_win <= 0.25.
- Candidate windows require >= 10 data points.
- Manuscript metrics are reported only if the selected window has >= 5 peaks and >= 5
  valleys/troughs.
- Peaks/troughs are detected from the unsmoothed, baseline-corrected force trace.
- Peak prominence = 0.003 * global maximum corrected force.

This module reports protocol-defined descriptors: Fc, Fci, Fc/w, Fci/w, PSI, SSA,
SSA_norm, selected-window CV, displacement/break proxy, and group repeatability.
It does not compute IC-Peel, Kendall, Gc, or Gci values.

Prepared for the Adhesion Paper Drafting workflow.

Version 1.4.0-rc15 adds window-feasibility warnings, cleaner validation separation, and stronger profile/benchmark guardrails,
metric-quality flags, suggested global settings, and improved QC caution banners.
"""

from __future__ import annotations

import os
import re
import gc
import json
import math
import time
import zipfile
import hashlib
import warnings
from dataclasses import dataclass, asdict, field
from datetime import datetime
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from scipy.signal import find_peaks

warnings.filterwarnings("ignore")

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None

# =============================================================================
# Constants
# =============================================================================

LBF_TO_N = 4.4482216152605
IN_TO_MM = 25.4

SUPPORTED_FORCE_UNITS = {"auto", "N", "n", "mN", "mn", "lbf"}
SUPPORTED_DISPLACEMENT_UNITS = {"auto", "mm", "um", "µm", "μm", "cm", "in"}

# In-session cache populated during scan_workbook(). This prevents re-reading the
# same large Excel sheets immediately during analysis and avoids fragile repeated
# workbook reads in some notebook runtimes.
_WORKBOOK_DATA_CACHE: Dict[Tuple[str, str], pd.DataFrame] = {}


__version__ = "1.4.0-rc15"

# Color-blind-safe, paper-matched palette used consistently in QC plots and guide figures.
PAPER_COLORS = {
    "E6000": "#1177B2",
    "Fabri-Fuse": "#5AA7D6",
    "FABRI-FUSE": "#5AA7D6",
    "JB Weld": "#8C8C8C",
    "J-B Weld": "#8C8C8C",
    "Elephant 407": "#DD6B00",
    "407": "#DD6B00",
    "SCOTCH": "black",
    "Without Glue": "black",
    "With Glue (Fabri-Fuse)": "#5AA7D6",
    "With E6000": "#1177B2",
}

QC_COLORS = {
    "force": "black",
    "selected_window": "#D7ECF7",
    "peak": "#CC3377",       # high-contrast magenta
    "trough": "#009E73",     # color-blind-safe green
    "fc": "#1177B2",
    "fci": "#CC3377",
    "fmean": "#DD6B00",
    "window_edge": "#222222",
    "drop": "#DD6B00",
    "outside": "#CFCFCF",
}

PLOT_STYLE = {
    "font.family": "sans-serif",
    "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
    "font.size": 14,
    "font.weight": "bold",
    "axes.labelweight": "bold",
    "axes.titleweight": "bold",
    "axes.linewidth": 1.8,
    "xtick.major.width": 1.8,
    "ytick.major.width": 1.8,
    "xtick.direction": "out",
    "ytick.direction": "in",
    "legend.frameon": False,
    "figure.dpi": 300,
    "savefig.dpi": 600,
    "savefig.facecolor": "white",
    "figure.facecolor": "white",
    "axes.facecolor": "white",
    "errorbar.capsize": 4,
}

LOCKED_MANUSCRIPT_PROFILE = {
    "start_offset_mm": 5.0,
    "window_length_mm": 25.0,
    "tail_exclude_frac": 0.20,
    "max_window_drift_frac": 0.25,
    "peak_prominence_frac": 0.003,
    "min_peak_distance_points": 1,
    "min_window_points": 10,
    "required_top_peaks": 5,
    "required_bottom_troughs": 5,
    "baseline_method": "first",
    "clip_force_negative_to_zero": True,
    "drop_zero_force_rows": True,
    "preserve_original_row_order": True,
    "drop_threshold_frac": 0.10,
    "break_proxy_consecutive_points": 1,
}


def apply_paper_plot_style() -> None:
    """Apply the paper-matched plotting style for QC and summary figures."""
    plt.rcParams.update(PLOT_STYLE)


def get_color_for_label(label: str, default: str = "#1177B2") -> str:
    """Return a stable paper-palette color for an adhesive/group label."""
    text = str(label or "")
    for key, color in PAPER_COLORS.items():
        if key.lower() in text.lower():
            return color
    return default


def method_profile_hash(profile: Optional[Dict[str, Any]] = None) -> str:
    """Stable SHA-256 hash of the locked manuscript profile."""
    payload = profile if profile is not None else LOCKED_MANUSCRIPT_PROFILE
    encoded = json.dumps(payload, sort_keys=True, default=str).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def validate_against_locked_manuscript_profile(config: "ManuscriptConfig") -> Dict[str, Any]:
    """Check whether analysis-sensitive parameters match the locked manuscript baseline."""
    changed = []
    for key, expected in LOCKED_MANUSCRIPT_PROFILE.items():
        actual = getattr(config, key, None)
        if isinstance(expected, float):
            same = math.isclose(float(actual), float(expected), rel_tol=0, abs_tol=1e-12)
        else:
            same = actual == expected
        if not same:
            changed.append({"parameter": key, "expected": expected, "actual": actual})
    return {
        "is_manuscript_baseline": len(changed) == 0,
        "changed_parameters": changed,
        "method_profile_hash": method_profile_hash(),
    }


def metric_quality_flag_for_result(result: Dict[str, Any]) -> str:
    """Return a compact quality flag for a trace result.

    PASS/FAIL remains the metric-computation status. This flag is an interpretation
    layer for user review, especially when Top5/Bot5 points are spatially clustered.
    """
    status = str(result.get("status", ""))
    if status != "PASS":
        return "METRIC_FAIL"
    trust = str(result.get("extrema_trust_flag", ""))
    warnings_text = str(result.get("warnings", ""))
    if trust == "LOW_EXTREMA_CLUSTERED" or "Extrema clustering warning" in warnings_text:
        return "PASS_WITH_CAUTION_LOW_EXTREMA_CLUSTERED"
    if str(result.get("failure_reason", "")):
        return "REVIEW_REQUIRED"
    return "QUALITY_OK"


def manual_review_required_for_result(result: Dict[str, Any]) -> bool:
    """Return True when a trace should be manually inspected before interpretation."""
    return metric_quality_flag_for_result(result) not in {"QUALITY_OK"}


def build_manuscript_baseline_config(
    input_path: str,
    output_root: str = "Peel_Analysis_Output_v1_4_0_rc15",
    width_mm: float = 20.0,
    thickness_mm: float = 0.40,
    bond_length_mm: float = 25.0,
    force_unit: str = "N",
    displacement_unit: str = "mm",
    assume_global_units: bool = True,
    export_qc_plots: bool = True,
    export_group_summary: bool = True,
    export_excel: bool = True,
    make_zip: bool = True,
) -> "ManuscriptConfig":
    """Create a manuscript-baseline config while keeping sensitive analysis settings locked."""
    cfg = ManuscriptConfig(
        input_path=input_path,
        output_root=output_root,
        width_mm=width_mm,
        thickness_mm=thickness_mm,
        bond_length_mm=bond_length_mm,
        assume_global_units=assume_global_units,
        force_unit=force_unit,
        displacement_unit=displacement_unit,
        export_qc_plots=export_qc_plots,
        export_group_summary=export_group_summary,
        export_excel=export_excel,
        make_zip=make_zip,
    )
    for key, value in LOCKED_MANUSCRIPT_PROFILE.items():
        setattr(cfg, key, value)
    cfg.method_profile = "manuscript_baseline_v1"
    cfg.allow_non_manuscript_run = False
    return cfg

TRUE_STRINGS = {"true", "t", "yes", "y", "1", "include", "included"}
FALSE_STRINGS = {"false", "f", "no", "n", "0", "exclude", "excluded"}

FORCE_HEADER_TERMS = [
    "force", "load", "peel force", "load cell", "loadcell", "tensile load", "peak load"
]
DISP_HEADER_TERMS = [
    "distance", "displacement", "extension", "crosshead", "cross-head", "stroke", "travel", "position", "elongation", "disp"
]
TIME_HEADER_TERMS = ["time", "elapsed", "seconds", "sec", "test time"]

FORCE_UNIT_PATTERNS = [
    (re.compile(r"(?:\(|\[|_|\b)(mn)(?:\)|\]|_|\b)", re.I), "mN"),
    (re.compile(r"(?:\(|\[|_|\b)(lbf|lb[f]?|pound[- ]?force)(?:\)|\]|_|\b)", re.I), "lbf"),
    (re.compile(r"(?:\(|\[|_|\b)(n|newton|newtons)(?:\)|\]|_|\b)", re.I), "N"),
]
DISP_UNIT_PATTERNS = [
    (re.compile(r"(?:\(|\[|_|\b)(mm|millimeter|millimeters)(?:\)|\]|_|\b)", re.I), "mm"),
    (re.compile(r"(?:\(|\[|_|\b)(um|µm|μm|micron|microns|micrometer|micrometers)(?:\)|\]|_|\b)", re.I), "um"),
    (re.compile(r"(?:\(|\[|_|\b)(cm|centimeter|centimeters)(?:\)|\]|_|\b)", re.I), "cm"),
    (re.compile(r"(?:\(|\[|_|\b)(in|inch|inches)(?:\)|\]|_|\b)", re.I), "in"),
]

# =============================================================================
# Dataclasses
# =============================================================================

@dataclass
class ManuscriptConfig:
    """Configuration for manuscript-guided paper-replica analysis."""

    input_path: str
    output_root: str = "Peel_Analysis_Output_v1_4_0_rc15"

    # Paper/manuscript geometry defaults. Sheet mapping can override these.
    width_mm: float = 20.0
    thickness_mm: float = 0.40
    bond_length_mm: float = 25.0

    # Unit handling. If assume_global_units is True, these are used unless headers contradict them.
    assume_global_units: bool = True
    force_unit: str = "N"
    displacement_unit: str = "mm"
    stop_on_unit_contradiction: bool = True

    # Optional mapping files generated by scan_workbook().
    sheet_mapping_file: str = ""
    group_defaults_file: str = ""

    # Inclusion/exclusion.
    exclude_sheet_if_contains: Sequence[str] = field(default_factory=list)

    # Manuscript window and peak extraction parameters.
    start_offset_mm: float = 5.0
    window_length_mm: float = 25.0
    tail_exclude_frac: float = 0.20
    max_window_drift_frac: float = 0.25
    peak_prominence_frac: float = 0.003
    min_peak_distance_points: int = 1
    min_window_points: int = 10
    required_top_peaks: int = 5
    required_bottom_troughs: int = 5

    # Preprocessing.
    baseline_method: str = "first"
    clip_force_negative_to_zero: bool = True
    drop_zero_force_rows: bool = True
    preserve_original_row_order: bool = True
    drop_threshold_frac: float = 0.10
    break_proxy_consecutive_points: int = 1

    # Extrema spread / clumping check. Baseline mode reports trust flags; strict QC can fail.
    extrema_spread_warning_frac: float = 0.10
    strict_extrema_spread_check: bool = False

    # Output options.
    export_qc_plots: bool = True
    export_group_summary: bool = True
    export_excel: bool = True
    export_column_audit: bool = True
    make_zip: bool = True
    dpi: int = 220

    # Method-profile integrity. Geometry and units are user-specific; analysis-sensitive
    # parameters should remain locked for manuscript-guided baseline runs.
    method_profile: str = "manuscript_baseline_v1"
    allow_non_manuscript_run: bool = False

    # Validation options.
    run_scotch_validation_if_detected: bool = True
    validation_relative_tolerance: float = 0.015  # 1.5% for row-level metrics
    validation_absolute_tolerance: float = 0.015


@dataclass
class ColumnCandidate:
    column: str
    force_score: float
    displacement_score: float
    force_unit: str = ""
    displacement_unit: str = ""
    reason: str = ""


@dataclass
class ColumnAssignment:
    force_column: str
    displacement_column: str
    confidence: str
    force_unit_detected: str
    displacement_unit_detected: str
    pair_score: float
    reason: str
    needs_user_confirmation: bool
    audit_rows: List[Dict[str, Any]]


@dataclass
class TraceInput:
    sheet: str
    main_group: str
    adhesive: str
    replicate_id: str
    include: bool
    width_mm: float
    thickness_mm: float
    bond_length_mm: float
    force_column: str
    displacement_column: str
    force_unit: str
    displacement_unit: str
    data: pd.DataFrame
    column_confidence: str
    column_reason: str
    warnings: List[str] = field(default_factory=list)

# =============================================================================
# Utility helpers
# =============================================================================

def now_stamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def ensure_dir(path: str) -> str:
    os.makedirs(path, exist_ok=True)
    return path


def safe_filename(text: str, max_len: int = 150) -> str:
    text = re.sub(r"[^A-Za-z0-9_\-. ]+", "_", str(text))
    text = re.sub(r"\s+", "_", text).strip("_")
    return (text or "unnamed")[:max_len]


def sha256_file(path: str, block_size: int = 1024 * 1024) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(block_size), b""):
            h.update(block)
    return h.hexdigest()


def normalize_header(x: Any) -> str:
    return re.sub(r"\s+", " ", str(x).strip())


def header_score(header: str, terms: Sequence[str]) -> float:
    h = header.lower()
    score = 0.0
    for term in terms:
        if term in h:
            score += 3.0 if len(term) > 4 else 2.0
    return score


def detect_unit_from_header(header: str, kind: str) -> str:
    patterns = FORCE_UNIT_PATTERNS if kind == "force" else DISP_UNIT_PATTERNS
    for pattern, unit in patterns:
        if pattern.search(header):
            return unit
    return ""


def numeric_array(series: pd.Series) -> np.ndarray:
    return pd.to_numeric(series, errors="coerce").to_numpy(dtype=float)


def numeric_quality(arr: np.ndarray) -> Dict[str, float]:
    mask = np.isfinite(arr)
    n = int(mask.sum())
    if n < 2:
        return {"n": n, "range": 0.0, "monotonic_fraction": 0.0, "nonmonotonic_fraction": 0.0}
    v = arr[mask]
    diffs = np.diff(v)
    nonzero = diffs[np.abs(diffs) > 1e-12]
    if nonzero.size == 0:
        mono = 0.0
    else:
        pos_frac = float(np.mean(nonzero > 0))
        neg_frac = float(np.mean(nonzero < 0))
        mono = max(pos_frac, neg_frac)
    return {
        "n": n,
        "range": float(np.nanmax(v) - np.nanmin(v)) if n else 0.0,
        "monotonic_fraction": mono,
        "nonmonotonic_fraction": 1.0 - mono,
    }


def unit_convert_force(values: np.ndarray, unit: str) -> np.ndarray:
    u = (unit or "N").strip()
    if u in {"N", "n", "auto", ""}:
        return values.astype(float)
    if u in {"mN", "mn"}:
        return values.astype(float) / 1000.0
    if u == "lbf":
        return values.astype(float) * LBF_TO_N
    raise ValueError(f"Unsupported force unit: {unit}")


def unit_convert_displacement(values: np.ndarray, unit: str) -> np.ndarray:
    u = (unit or "mm").strip()
    if u in {"mm", "auto", ""}:
        return values.astype(float)
    if u in {"um", "µm", "μm"}:
        return values.astype(float) / 1000.0
    if u == "cm":
        return values.astype(float) * 10.0
    if u == "in":
        return values.astype(float) * IN_TO_MM
    raise ValueError(f"Unsupported displacement unit: {unit}")


def parse_bool(value: Any, default: bool = True) -> bool:
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return default
    s = str(value).strip().lower()
    if s in TRUE_STRINGS:
        return True
    if s in FALSE_STRINGS:
        return False
    return default


def infer_group_and_adhesive(sheet: str) -> Tuple[str, str, str]:
    """Lightweight inference for group/adhesive/replicate labels from a sheet name.

    Users can override these labels in sheet_mapping_review.csv.
    """
    raw = str(sheet)
    low = raw.lower()
    group = "UNKNOWN"
    adhesive = "UNKNOWN"
    if "scotch" in low:
        group = "SCOTCH TAPE T PEEL"
        adhesive = "SCOTCH"
    elif "82" in low and "n" in low:
        group = "82N"
    elif "88" in low and "n" in low:
        group = "88N"
    parts = re.split(r"[_\- ]+", raw)
    known = ["e6000", "fabri", "fuse", "jb", "weld", "407", "elephant", "scotch", "tape"]
    hit = [p for p in parts if p.lower() in known]
    if adhesive == "UNKNOWN" and hit:
        adhesive = " ".join(hit).upper().replace("FABRI FUSE", "FABRI-FUSE")
    rep = raw
    return group, adhesive, rep

# =============================================================================
# Reading workbooks and detecting columns
# =============================================================================

def read_table_from_excel_sheet(path: str, sheet: str, header_row: int = 1) -> pd.DataFrame:
    """Read one Excel sheet efficiently and normalize headers.

    Pandas/openpyxl is faster than manual row iteration for many instrument
    workbooks with tens of thousands of rows. A read-only openpyxl fallback is
    retained for unusual files.
    """
    try:
        df = pd.read_excel(path, sheet_name=sheet, engine="openpyxl", header=header_row - 1)
        df.columns = [normalize_header(c) for c in df.columns]
        # Drop columns that are entirely blank or unnamed after normalization.
        keep = [c for c in df.columns if c and not str(c).lower().startswith("unnamed")]
        if keep:
            df = df.loc[:, keep]
        return df
    except Exception:
        if load_workbook is None:
            raise ImportError("openpyxl is required to read .xlsx files")
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet]
        rows = ws.iter_rows(values_only=True)
        header = None
        for idx, row in enumerate(rows, start=1):
            if idx == header_row:
                header = [normalize_header(x) for x in row]
                break
        if header is None:
            wb.close()
            return pd.DataFrame()
        last = 0
        for i, h in enumerate(header):
            if h and h.lower() != "none":
                last = i + 1
        header = header[:last]
        data = []
        for row in rows:
            data.append(tuple(row[:last]))
        wb.close()
        return pd.DataFrame(data, columns=header)


def is_user_data_file(path: str) -> bool:
    """Return True only for plausible user input data files.

    The notebook intentionally does not accept .txt as input because the manuscript
    workflow expects structured Excel/CSV tables with explicit columns. Package files,
    templates, expected-output validation tables, and prior analysis outputs are excluded
    so they are not accidentally selected as peel data.
    """
    p = os.path.abspath(str(path))
    name = os.path.basename(p)
    lname = name.lower()
    ext = os.path.splitext(lname)[1]
    if ext not in {".xlsx", ".xls", ".xlsm", ".csv"}:
        return False
    if name.startswith("~$"):
        return False
    exclude_exact = {
        # Included package validation workbook. It is run by the built-in validation cell,
        # not offered as a normal user-upload candidate.
        "scothtapetpeel.xlsx",
        "scothtapetpeel.xls",
        "scothtapetpeel.csv",
        "scotchtapetpeel.xlsx",
        "scotchtapetpeel.xls",
        "scotchtapetpeel.csv",
    }
    if lname in exclude_exact:
        return False
    exclude_tokens = [
        "template", "expected_scotch", "method_profile", "user_method_profile",
        "user_benchmark", "benchmark_expected", "benchmark_metadata", "group_defaults",
        "requirements", "readme", "citation", "license",
        "paper_metrics", "group_summary", "formula_dictionary", "metric_guide",
        "output_file_guide", "input_audit", "readiness", "validation",
        "column_assignment", "advanced_column", "provenance", "analysis_report",
        "manuscript_baseline_outputs", "scotch_validation", "output_consistency",
        "diagnostic_summary", "issue_action", "trace_integrity", "run_history",
    ]
    if any(tok in lname for tok in exclude_tokens):
        return False
    if "peel_analysis_output" in p.lower():
        return False
    return True


def list_input_files(directory: str = ".") -> List[str]:
    files = []
    for name in sorted(os.listdir(directory)):
        path = os.path.join(directory, name)
        if os.path.isfile(path) and is_user_data_file(path):
            files.append(os.path.abspath(path))
    return files


def workbook_sheet_names(path: str) -> List[str]:
    ext = os.path.splitext(path)[1].lower()
    if ext in {".xlsx", ".xlsm", ".xls"}:
        if load_workbook is None:
            raise ImportError("openpyxl is required to read Excel files")
        wb = load_workbook(path, read_only=True, data_only=True)
        names = list(wb.sheetnames)
        wb.close()
        return names
    return [os.path.basename(path)]




def validate_input_file(path: str) -> Dict[str, Any]:
    """Check that a user-selected input file exists, has a supported extension, and is readable.

    This prevents non-data files (for example requirements files) from being accepted as peel
    data and failing later with obscure zip/Excel parsing errors.
    """
    info = {
        "path": path,
        "exists": False,
        "supported_extension": False,
        "readable": False,
        "extension": os.path.splitext(str(path))[1].lower(),
        "n_sheets": 0,
        "sheets": [],
        "message": "",
        "sha256": "",
    }
    if not path:
        info["message"] = "No file path was provided."
        return info
    if not os.path.exists(path):
        info["message"] = f"File not found: {path}"
        return info
    info["exists"] = True
    ext = info["extension"]
    if ext not in {".xlsx", ".xls", ".xlsm", ".csv"}:
        info["message"] = f"Unsupported file extension '{ext}'. Use .xlsx, .xls, .xlsm, or .csv."
        return info
    info["supported_extension"] = True
    try:
        if ext in {".xlsx", ".xls", ".xlsm"}:
            if load_workbook is None:
                raise ImportError("openpyxl is required to read Excel files")
            wb = load_workbook(path, read_only=True, data_only=True)
            info["sheets"] = list(wb.sheetnames)
            wb.close()
            info["n_sheets"] = len(info["sheets"])
            if info["n_sheets"] == 0:
                raise ValueError("Excel workbook contains no sheets")
        else:
            preview = pd.read_csv(path, nrows=5)
            info["sheets"] = [os.path.basename(path)]
            info["n_sheets"] = 1
            if preview.shape[1] < 2:
                raise ValueError("CSV file must contain at least two columns")
        info["sha256"] = sha256_file(path)
        info["readable"] = True
        info["message"] = "File is readable."
    except zipfile.BadZipFile:
        info["message"] = "The file has an Excel extension but is not a readable Excel workbook (BadZipFile). Re-upload or choose a valid .xlsx/.xlsm file."
    except Exception as exc:
        info["message"] = f"File readability check failed: {exc}"
    return info


def read_sheet_preview(path: str, sheet: str, nrows: int = 500) -> pd.DataFrame:
    ext = os.path.splitext(path)[1].lower()
    if ext in {".xlsx", ".xlsm", ".xls"}:
        df = read_table_from_excel_sheet(path, sheet, header_row=1)
        return df.head(nrows)
    return pd.read_csv(path, nrows=nrows)


def detect_force_displacement_columns(df: pd.DataFrame) -> ColumnAssignment:
    """Detect force/displacement columns using header, unit, numeric behavior, and pairwise plausibility."""
    candidates: List[ColumnCandidate] = []
    for col in df.columns:
        h = normalize_header(col)
        arr = numeric_array(df[col])
        q = numeric_quality(arr)
        f_unit = detect_unit_from_header(h, "force")
        x_unit = detect_unit_from_header(h, "displacement")
        f_score = header_score(h, FORCE_HEADER_TERMS) + (4.0 if f_unit else 0.0)
        x_score = header_score(h, DISP_HEADER_TERMS) + (4.0 if x_unit else 0.0)
        if q["n"] >= 10 and q["range"] > 0:
            f_score += 1.0
            x_score += 1.0
        # Displacement is usually monotonic in row order; force is usually not strictly monotonic.
        if q["monotonic_fraction"] >= 0.92:
            x_score += 2.0
        if q["nonmonotonic_fraction"] >= 0.05:
            f_score += 1.0
        candidates.append(ColumnCandidate(h, f_score, x_score, f_unit, x_unit, reason=f"n={q['n']}, range={q['range']:.4g}, monotonic={q['monotonic_fraction']:.3f}"))

    audit_rows: List[Dict[str, Any]] = []
    best_pair = None
    headers = [normalize_header(c) for c in df.columns]
    for f_cand in candidates:
        for x_cand in candidates:
            if f_cand.column == x_cand.column:
                continue
            f_arr = numeric_array(df[f_cand.column])
            x_arr = numeric_array(df[x_cand.column])
            mask = np.isfinite(f_arr) & np.isfinite(x_arr)
            pair_score = f_cand.force_score + x_cand.displacement_score
            if mask.sum() >= 10:
                fq = numeric_quality(f_arr[mask])
                xq = numeric_quality(x_arr[mask])
                if xq["monotonic_fraction"] >= 0.90:
                    pair_score += 3.0
                if fq["range"] > 0:
                    pair_score += 1.0
                # Penalize suspiciously identical columns.
                try:
                    corr = np.corrcoef(f_arr[mask], x_arr[mask])[0, 1]
                    if np.isfinite(corr) and abs(corr) > 0.995:
                        pair_score -= 1.0
                except Exception:
                    corr = np.nan
            else:
                corr = np.nan
            row = {
                "candidate_force_column": f_cand.column,
                "candidate_displacement_column": x_cand.column,
                "force_score": round(f_cand.force_score, 3),
                "displacement_score": round(x_cand.displacement_score, 3),
                "pair_score": round(pair_score, 3),
                "force_unit_detected": f_cand.force_unit,
                "displacement_unit_detected": x_cand.displacement_unit,
                "reason": f"force: {f_cand.reason}; displacement: {x_cand.reason}",
            }
            audit_rows.append(row)
            if best_pair is None or pair_score > best_pair[0]:
                best_pair = (pair_score, f_cand, x_cand, row)

    if best_pair is None:
        return ColumnAssignment("", "", "LOW", "", "", 0.0, "No valid numeric force/displacement pair found.", True, audit_rows)

    sorted_scores = sorted([r["pair_score"] for r in audit_rows], reverse=True)
    margin = sorted_scores[0] - sorted_scores[1] if len(sorted_scores) > 1 else sorted_scores[0]
    score, f_cand, x_cand, _ = best_pair
    if score >= 12 and margin >= 2.0:
        conf = "HIGH"
    elif score >= 8 and margin >= 0.75:
        conf = "MEDIUM"
    else:
        conf = "LOW"
    reason = (
        f"Selected force='{f_cand.column}', displacement='{x_cand.column}'. "
        f"Pair score={score:.2f}, margin={margin:.2f}. "
        f"Force evidence: score={f_cand.force_score:.2f}, unit={f_cand.force_unit or 'not detected'}. "
        f"Displacement evidence: score={x_cand.displacement_score:.2f}, unit={x_cand.displacement_unit or 'not detected'}."
    )
    return ColumnAssignment(
        force_column=f_cand.column,
        displacement_column=x_cand.column,
        confidence=conf,
        force_unit_detected=f_cand.force_unit,
        displacement_unit_detected=x_cand.displacement_unit,
        pair_score=float(score),
        reason=reason,
        needs_user_confirmation=(conf == "LOW"),
        audit_rows=audit_rows,
    )


def unit_resolution(config: ManuscriptConfig, assignment: ColumnAssignment) -> Tuple[str, str, List[str], bool]:
    warnings_out = []
    contradiction = False
    force_unit = assignment.force_unit_detected or (config.force_unit if config.force_unit != "auto" else "")
    disp_unit = assignment.displacement_unit_detected or (config.displacement_unit if config.displacement_unit != "auto" else "")
    if config.assume_global_units:
        if assignment.force_unit_detected and config.force_unit != "auto" and assignment.force_unit_detected.lower() != config.force_unit.lower():
            contradiction = True
            warnings_out.append(f"Header suggests force unit {assignment.force_unit_detected}, but global force unit is {config.force_unit}.")
        if assignment.displacement_unit_detected and config.displacement_unit != "auto" and assignment.displacement_unit_detected.lower() != config.displacement_unit.lower():
            contradiction = True
            warnings_out.append(f"Header suggests displacement unit {assignment.displacement_unit_detected}, but global displacement unit is {config.displacement_unit}.")
        force_unit = config.force_unit if config.force_unit != "auto" else (assignment.force_unit_detected or "")
        disp_unit = config.displacement_unit if config.displacement_unit != "auto" else (assignment.displacement_unit_detected or "")
    if not force_unit:
        warnings_out.append("Force unit unresolved.")
    if not disp_unit:
        warnings_out.append("Displacement unit unresolved.")
    return force_unit, disp_unit, warnings_out, contradiction

# =============================================================================
# Preprocessing and analysis
# =============================================================================

def preprocess_trace(df: pd.DataFrame, force_col: str, displacement_col: str, force_unit: str, displacement_unit: str, config: ManuscriptConfig) -> Tuple[np.ndarray, np.ndarray, Dict[str, Any]]:
    f_raw = numeric_array(df[force_col])
    x_raw = numeric_array(df[displacement_col])
    mask = np.isfinite(f_raw) & np.isfinite(x_raw)
    f_raw = f_raw[mask]
    x_raw = x_raw[mask]
    f = unit_convert_force(f_raw, force_unit)
    x = unit_convert_displacement(x_raw, displacement_unit)

    displacement_flipped = False
    if len(x) >= 3:
        dx = np.diff(x)
        nonzero_dx = dx[np.abs(dx) > 1e-12]
        median_dx = np.nanmedian(nonzero_dx) if nonzero_dx.size else 0
        if median_dx < 0 or abs(np.nanmin(x)) > max(abs(np.nanmax(x)), 1e-12) * 2:
            x = -x
            displacement_flipped = True

    if config.baseline_method == "first":
        baseline = f[0] if len(f) else 0.0
    elif config.baseline_method == "median5":
        baseline = float(np.nanmedian(f[:5])) if len(f) else 0.0
    else:
        baseline = f[0] if len(f) else 0.0
    f = f - baseline
    if config.clip_force_negative_to_zero:
        f = np.where(f < 0, 0.0, f)
    if config.drop_zero_force_rows:
        keep = f > 0
        f = f[keep]
        x = x[keep]

    # Preserve row order. Window search requires nondecreasing displacement; if duplicated/reversed local points exist, sort only when necessary.
    sorted_for_analysis = False
    if len(x) >= 2 and np.nanmean(np.diff(x) >= -1e-9) < 0.98:
        order = np.argsort(x)
        x = x[order]
        f = f[order]
        sorted_for_analysis = True

    info = {
        "raw_numeric_points": int(mask.sum()),
        "processed_points": int(len(f)),
        "force_baseline_N": float(baseline),
        "displacement_flipped": bool(displacement_flipped),
        "sorted_for_analysis": bool(sorted_for_analysis),
    }
    return x, f, info


def cumulative_window_search(x: np.ndarray, f: np.ndarray, config: ManuscriptConfig) -> Tuple[Optional[Dict[str, Any]], Optional[Tuple[int, int]]]:
    if len(x) < config.min_window_points:
        return None, None
    x_min = float(np.nanmin(x))
    x_max = float(np.nanmax(x))
    span = x_max - x_min
    if not np.isfinite(span) or span <= 0:
        return None, None
    tail_limit = x_min + (1.0 - config.tail_exclude_frac) * span
    L = float(config.window_length_mm)

    c_y = np.r_[0.0, np.cumsum(f)]
    c_y2 = np.r_[0.0, np.cumsum(f * f)]
    c_x = np.r_[0.0, np.cumsum(x)]
    c_x2 = np.r_[0.0, np.cumsum(x * x)]
    c_xy = np.r_[0.0, np.cumsum(x * f)]

    best = None
    best_ij = None
    for i, xi in enumerate(x):
        if xi < x_min + config.start_offset_mm:
            continue
        xe = xi + L
        if xe > tail_limit + 1e-9:
            break
        j = int(np.searchsorted(x, xe, side="right"))
        n = j - i
        if n < config.min_window_points:
            continue
        sy = c_y[j] - c_y[i]
        sy2 = c_y2[j] - c_y2[i]
        sx = c_x[j] - c_x[i]
        sx2 = c_x2[j] - c_x2[i]
        sxy = c_xy[j] - c_xy[i]
        mean = sy / n
        if mean <= 0 or not np.isfinite(mean):
            continue
        var = (sy2 - sy * sy / n) / (n - 1) if n > 1 else 0.0
        sd = math.sqrt(max(float(var), 0.0))
        denom = sx2 - sx * sx / n
        slope = (sxy - sx * sy / n) / denom if abs(denom) > 1e-12 else 0.0
        drift_ratio = abs(slope) * L / mean
        if drift_ratio > config.max_window_drift_frac:
            continue
        if best is None or sd < best["fstd_N"] - 1e-12:
            best = {
                "window_start_mm": float(xi),
                "window_end_mm": float(xe),
                "window_points": int(n),
                "fmean_N": float(mean),
                "fstd_N": float(sd),
                "window_cv_pct": float(100 * sd / mean) if mean else np.nan,
                "window_slope_N_per_mm": float(slope),
                "window_drift_ratio": float(drift_ratio),
                "peel_window_travel_mm": float(L),
                "tail_limit_mm": float(tail_limit),
            }
            best_ij = (i, j)
    return best, best_ij


def detect_break_proxy(x: np.ndarray, f: np.ndarray, threshold_frac: float = 0.10) -> float:
    if len(x) == 0 or len(f) == 0:
        return np.nan
    fmax = float(np.nanmax(f))
    if not np.isfinite(fmax) or fmax <= 0:
        return float(np.nanmax(x))
    imax = int(np.nanargmax(f))
    threshold = threshold_frac * fmax
    after = np.where((np.arange(len(f)) > imax) & (f <= threshold))[0]
    if after.size:
        return float(x[after[0]])
    return float(np.nanmax(x))


def robust_noise_floor(f: np.ndarray, n_first: int = 25) -> float:
    if len(f) < 5:
        return np.nan
    seg = f[: min(n_first, len(f))]
    med = np.nanmedian(seg)
    mad = np.nanmedian(np.abs(seg - med))
    return float(1.4826 * mad)


def detect_break_proxy_after_window(
    x: np.ndarray,
    f: np.ndarray,
    fc: float,
    window_end_index: int,
    threshold_frac: float = 0.10,
    consecutive_points: int = 3,
) -> float:
    """Return post-window drop/break proxy based on a fraction of Fc.

    This avoids using a large global maximum or an initial grip/slack peak as the
    drop reference. A drop is accepted only after the selected analysis window and
    only when the signal remains below threshold for a short sustained run.
    """
    if len(x) == 0 or len(f) == 0 or not np.isfinite(fc) or fc <= 0:
        return np.nan
    threshold = threshold_frac * float(fc)
    start = max(int(window_end_index), 0)
    if start >= len(f):
        return np.nan
    below = np.asarray(f[start:] <= threshold, dtype=bool)
    k = max(1, int(consecutive_points))
    if len(below) < k:
        return np.nan
    # First index where k consecutive points are below threshold.
    run = np.convolve(below.astype(int), np.ones(k, dtype=int), mode="valid")
    hits = np.where(run >= k)[0]
    if hits.size:
        return float(x[start + int(hits[0])])
    return np.nan


def analyze_trace(trace: TraceInput, config: ManuscriptConfig) -> Tuple[Dict[str, Any], Optional[pd.DataFrame]]:
    warnings_out = list(trace.warnings)
    base = {
        "sheet": trace.sheet,
        "main_group": trace.main_group,
        "adhesive": trace.adhesive,
        "replicate_id": trace.replicate_id,
        "method": "ASTM D2724-style manuscript mode (Top5+Bot5)",
        "w_mm": trace.width_mm,
        "t_mm": trace.thickness_mm,
        "bond_length_mm": trace.bond_length_mm,
        "force_column": trace.force_column,
        "displacement_column": trace.displacement_column,
        "force_unit": trace.force_unit,
        "displacement_unit": trace.displacement_unit,
        "column_confidence": trace.column_confidence,
    }
    try:
        x, f, prep = preprocess_trace(trace.data, trace.force_column, trace.displacement_column, trace.force_unit, trace.displacement_unit, config)
    except Exception as exc:
        base.update({"status": "FAIL", "failure_reason": f"Preprocessing failed: {exc}", "warnings": "; ".join(warnings_out)})
        return base, None

    base.update(prep)
    base["distance_max_mm"] = float(np.nanmax(x)) if len(x) else np.nan
    base["distance_min_mm"] = float(np.nanmin(x)) if len(x) else np.nan
    # Break/drop proxy is intentionally calculated after Fc is known so it uses
    # 0.10 * Fc and searches only after the selected analysis window. If a
    # post-window force drop is not observed, the final recorded displacement is
    # reported as a terminal-displacement fallback and explicitly flagged.
    base["drop_threshold_N"] = np.nan
    base["distance_at_drop_mm"] = np.nan
    base["displacement_at_break_mm"] = np.nan
    base["break_proxy_status"] = "NOT_EVALUATED"
    base["break_proxy_basis"] = "0.10 × Fc after selected window; fallback uses final recorded displacement"
    base["noise_floor_N"] = robust_noise_floor(f)

    if len(x) < config.min_window_points:
        base.update({
            "status": "FAIL",
            "failure_reason": f"Too few processed numeric points ({len(x)} < {config.min_window_points}).",
            "warnings": "; ".join(warnings_out),
        })
        base["displacement_at_break_mm"] = float(np.nanmax(x)) if len(x) else np.nan
        base["break_proxy_status"] = "METRIC_FAIL_USED_FINAL_RECORDED_DISPLACEMENT"
        return base, pd.DataFrame({"distance_mm": x, "force_N": f})

    window, ij = cumulative_window_search(x, f, config)
    if window is None or ij is None:
        base.update({
            "status": "FAIL",
            "failure_reason": "No valid 25.0 mm manuscript window satisfied point, tail, and drift constraints.",
            "warnings": "; ".join(warnings_out),
        })
        base["displacement_at_break_mm"] = float(np.nanmax(x)) if len(x) else np.nan
        base["break_proxy_status"] = "NO_VALID_WINDOW_USED_FINAL_RECORDED_DISPLACEMENT"
        return base, pd.DataFrame({"distance_mm": x, "force_N": f})

    i, j = ij
    xw = x[i:j]
    fw = f[i:j]
    prominence = config.peak_prominence_frac * float(np.nanmax(f)) if len(f) else np.nan
    peaks, _ = find_peaks(fw, prominence=prominence, distance=config.min_peak_distance_points)
    troughs, _ = find_peaks(-fw, prominence=prominence, distance=config.min_peak_distance_points)

    peak_count = int(len(peaks))
    trough_count = int(len(troughs))
    base.update(window)
    base.update({
        "peak_prominence_N": float(prominence),
        "peak_count": peak_count,
        "trough_count": trough_count,
        "Top_Points": min(peak_count, config.required_top_peaks),
        "Bot_Points": min(trough_count, config.required_bottom_troughs),
    })

    if peak_count < config.required_top_peaks or trough_count < config.required_bottom_troughs:
        reason = (
            f"Incomplete extrema for manuscript metrics: peaks={peak_count}, troughs={trough_count}; "
            f"required {config.required_top_peaks} peaks and {config.required_bottom_troughs} troughs."
        )
        base.update({
            "status": "FAIL",
            "failure_reason": reason,
            "warnings": "; ".join(warnings_out),
        })
        window_df = pd.DataFrame({"distance_mm": x, "force_N": f, "in_selected_window": False})
        window_df.loc[i:j-1, "in_selected_window"] = True
        base["displacement_at_break_mm"] = float(np.nanmax(x)) if len(x) else np.nan
        base["break_proxy_status"] = "INCOMPLETE_EXTREMA_USED_FINAL_RECORDED_DISPLACEMENT"
        return base, window_df

    # Highest five peaks and lowest five valleys/troughs.
    top_idx = peaks[np.argsort(fw[peaks])[-config.required_top_peaks:]]
    bot_idx = troughs[np.argsort(fw[troughs])[:config.required_bottom_troughs]]

    # Extrema-clumping / horizontal-spread diagnostic. This is reported as a trust
    # flag in manuscript baseline mode and can be made a hard fail in strict QC mode.
    top_x = xw[top_idx]
    bot_x = xw[bot_idx]
    peak_spread_mm = float(np.nanmax(top_x) - np.nanmin(top_x)) if len(top_x) else np.nan
    trough_spread_mm = float(np.nanmax(bot_x) - np.nanmin(bot_x)) if len(bot_x) else np.nan
    min_required_spread = float(config.extrema_spread_warning_frac) * float(config.window_length_mm)
    clustered = (
        np.isfinite(peak_spread_mm) and np.isfinite(trough_spread_mm) and
        (peak_spread_mm < min_required_spread or trough_spread_mm < min_required_spread)
    )
    extrema_trust_flag = "LOW_EXTREMA_CLUSTERED" if clustered else "OK_EXTREMA_SPREAD"
    if clustered:
        warnings_out.append(
            f"Extrema clustering warning: Top5 peaks span {peak_spread_mm:.2f} mm and "
            f"Bot5 troughs span {trough_spread_mm:.2f} mm; minimum recommended spread is "
            f"{min_required_spread:.2f} mm. Inspect QC plot before interpreting Top5/Bottom5 metrics."
        )
        if config.strict_extrema_spread_check:
            base.update({
                "status": "FAIL",
                "failure_reason": (
                    f"Extrema clumping failed strict QC: Top5 spread={peak_spread_mm:.2f} mm, "
                    f"Bot5 spread={trough_spread_mm:.2f} mm, required >= {min_required_spread:.2f} mm."
                ),
                "peak_spread_mm": peak_spread_mm,
                "trough_spread_mm": trough_spread_mm,
                "extrema_spread_min_required_mm": min_required_spread,
                "extrema_trust_flag": extrema_trust_flag,
                "warnings": "; ".join(warnings_out),
            })
            window_df = pd.DataFrame({"distance_mm": x, "force_N": f, "in_selected_window": False})
            window_df.loc[i:j-1, "in_selected_window"] = True
            base["displacement_at_break_mm"] = float(np.nanmax(x)) if len(x) else np.nan
            base["break_proxy_status"] = "STRICT_EXTREMA_QC_FAIL_USED_FINAL_RECORDED_DISPLACEMENT"
            return base, window_df

    top_vals = fw[top_idx]
    bot_vals = fw[bot_idx]
    fci = float(np.mean(top_vals))
    fca = float(np.mean(bot_vals))
    fc = float(np.mean(np.r_[top_vals, bot_vals]))
    ssa = float(fci - fca)
    drop_threshold_N = float(config.drop_threshold_frac * fc) if np.isfinite(fc) else np.nan
    d_drop = detect_break_proxy_after_window(
        x, f, fc, j, config.drop_threshold_frac, config.break_proxy_consecutive_points
    )
    base["drop_threshold_N"] = drop_threshold_N
    base["distance_at_drop_mm"] = d_drop
    if np.isfinite(d_drop):
        base["displacement_at_break_mm"] = d_drop
        base["break_proxy_status"] = "DROP_FOUND_AFTER_SELECTED_WINDOW"
    else:
        base["displacement_at_break_mm"] = float(np.nanmax(x)) if len(x) else np.nan
        base["break_proxy_status"] = "NO_DROP_FOUND_USED_FINAL_RECORDED_DISPLACEMENT"
    fmean = float(window["fmean_N"])
    fstd = float(window["fstd_N"])
    psi = float(fmean / fstd) if fstd > 0 else np.inf
    ssa_norm = float(ssa / fmean) if fmean > 0 else np.nan
    width = trace.width_mm

    base.update({
        "status": "PASS",
        "failure_reason": "",
        "fc_N": fc,
        "fci_N": fci,
        "fca_N": fca,
        "fc_over_w_N_per_mm": fc / width if width else np.nan,
        "fci_over_w_N_per_mm": fci / width if width else np.nan,
        "ssa_abs_N": ssa,
        "ssa_norm": ssa_norm,
        "psi_mu_over_sigma": psi,
        "peak_spread_mm": peak_spread_mm,
        "trough_spread_mm": trough_spread_mm,
        "extrema_spread_min_required_mm": min_required_spread,
        "extrema_trust_flag": extrema_trust_flag,
        "sigma_ci_MPa": fci / (trace.width_mm * trace.thickness_mm) if trace.width_mm and trace.thickness_mm else np.nan,
        "ssa_to_noise_ratio": ssa / base["noise_floor_N"] if base.get("noise_floor_N") and np.isfinite(base.get("noise_floor_N")) and base.get("noise_floor_N") > 0 else np.nan,
        "warnings": "; ".join(warnings_out),
    })

    window_df = pd.DataFrame({"distance_mm": x, "force_N": f, "in_selected_window": False, "selected_peak": False, "selected_trough": False})
    window_df.loc[i:j-1, "in_selected_window"] = True
    # Translate window-local selected indices to global positions.
    for idx in top_idx:
        window_df.loc[i + int(idx), "selected_peak"] = True
    for idx in bot_idx:
        window_df.loc[i + int(idx), "selected_trough"] = True
    return base, window_df

# =============================================================================
# Plotting
# =============================================================================


def make_qc_plot(result: Dict[str, Any], trace_df: pd.DataFrame, outpath: str, config: ManuscriptConfig, validation_status: str = "") -> None:
    """Create a clean, consistent QC plot for all traces.

    This version uses one shared layout for Scotch reference plots and user traces:
    a full-width status title, two panels, white-backed annotations, and a
    figure-level legend outside dense trace regions.
    """
    if trace_df is None or trace_df.empty or "distance_mm" not in trace_df:
        return
    apply_paper_plot_style()
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch

    x = trace_df["distance_mm"].to_numpy(float)
    f = trace_df["force_N"].to_numpy(float)
    inwin = trace_df.get("in_selected_window", pd.Series(False, index=trace_df.index)).to_numpy(bool)
    peaks = trace_df.get("selected_peak", pd.Series(False, index=trace_df.index)).to_numpy(bool)
    troughs = trace_df.get("selected_trough", pd.Series(False, index=trace_df.index)).to_numpy(bool)

    trace_color = QC_COLORS["force"]
    quality_flag = str(result.get("metric_quality_flag") or metric_quality_flag_for_result(result))
    review_required = bool(result.get("manual_review_required", manual_review_required_for_result(result)))
    metric_status = str(result.get("status", ""))
    if metric_status == "PASS" and review_required:
        status_text = "METRIC PASS | REVIEW REQUIRED"
        status_color = QC_COLORS["drop"]
        status_bg = "#fff8e1"
    elif metric_status == "PASS":
        status_text = "METRIC PASS | QUALITY OK"
        status_color = "#2e7d32"
        status_bg = "#f4fbf5"
    else:
        status_text = f"METRIC {metric_status} | DO NOT INTERPRET TOP5/BOTTOM5 METRICS"
        status_color = "#b71c1c"
        status_bg = "#ffebee"
    if validation_status:
        status_text += f" | {validation_status}"

    fig, axes = plt.subplots(1, 2, figsize=(19.0, 6.4), constrained_layout=False)
    fig.patch.set_facecolor("white")
    fig.patch.set_alpha(1.0)
    fig.subplots_adjust(left=0.065, right=0.985, top=0.80, bottom=0.20, wspace=0.19)

    title = f"{result.get('sheet','trace')} | {status_text}"
    fig.text(
        0.5, 0.965, title,
        ha="center", va="top", fontsize=14.0, fontweight="bold", color="black",
        bbox=dict(boxstyle="round,pad=0.34", facecolor="white", edgecolor=status_color, linewidth=1.3, alpha=0.99),
    )

    if metric_status != "PASS":
        caution = "Metric extraction failed. Review the raw trace and failure reason before using this result."
        fig.text(
            0.5, 0.890, caution,
            ha="center", va="center", fontsize=10.2, fontweight="bold", color="#7f0000",
            bbox=dict(boxstyle="round,pad=0.35", facecolor="#ffebee", edgecolor="#b71c1c", linewidth=1.25, alpha=0.99),
        )
    elif review_required:
        caution = "Review required: metrics were computed, but QC flags indicate manual inspection is needed."
        fig.text(
            0.5, 0.890, caution,
            ha="center", va="center", fontsize=10.2, fontweight="bold", color="#5d4037",
            bbox=dict(boxstyle="round,pad=0.35", facecolor="#fff8e1", edgecolor=QC_COLORS["drop"], linewidth=1.2, alpha=0.99),
        )

    for _ax in axes:
        _ax.set_facecolor("white")
        for spine in _ax.spines.values():
            spine.set_linewidth(1.65)

    # ----------------------------- full trace -----------------------------
    ax = axes[0]
    ax.plot(x, f, lw=1.95, color=trace_color, zorder=2)
    if np.any(inwin):
        xs = result.get("window_start_mm", np.nan)
        xe = result.get("window_end_mm", np.nan)
        ax.axvspan(xs, xe, alpha=0.55, color=QC_COLORS["selected_window"], zorder=0)
        ax.axvline(xs, ls="--", lw=1.45, color=QC_COLORS["window_edge"], zorder=3)
        ax.axvline(xe, ls="--", lw=1.45, color=QC_COLORS["window_edge"], zorder=3)
    if np.isfinite(result.get("distance_at_drop_mm", np.nan)):
        ax.axvline(result.get("distance_at_drop_mm"), ls=":", lw=2.0, color=QC_COLORS["drop"], zorder=3)

    for key, label, style, color in [
        ("fmean_N", "Fmean", "--", QC_COLORS["fmean"]),
        ("fc_N", "Fc", "-.", QC_COLORS["fc"]),
        ("fci_N", "Fci", ":", QC_COLORS["fci"]),
    ]:
        val = result.get(key, np.nan)
        if np.isfinite(val):
            ax.axhline(val, ls=style, lw=1.9, color=color, zorder=1)

    ax.set_title("Full force trace", fontsize=12, fontweight="bold", pad=8)
    ax.set_xlabel("Displacement (mm)")
    ax.set_ylabel("Force (N)")
    ax.grid(True, alpha=0.15)
    summary = (
        f"Top={result.get('Top_Points','')}, Bot={result.get('Bot_Points','')} | "
        f"Peaks={result.get('peak_count','')}, Troughs={result.get('trough_count','')}\n"
        f"CV={result.get('window_cv_pct', np.nan):.1f}%, PSI={result.get('psi_mu_over_sigma', np.nan):.2f}, "
        f"SSA={result.get('ssa_abs_N', np.nan):.3g} N\n"
        f"Drift={result.get('window_drift_ratio', np.nan):.3f}; slope={result.get('window_slope_N_per_mm', np.nan):.3g} N/mm\n"
        f"Quality: {quality_flag}"
    )
    ax.text(
        0.025, 0.035, summary, transform=ax.transAxes, fontsize=8.5, ha="left", va="bottom",
        bbox=dict(boxstyle="round,pad=0.35", facecolor="white", alpha=0.98, edgecolor="0.45", linewidth=1.0),
    )

    # ----------------------------- zoom panel -----------------------------
    ax = axes[1]
    if np.any(inwin):
        ax.plot(x[~inwin], f[~inwin], lw=1.45, alpha=0.32, color=QC_COLORS["outside"])
        ax.plot(x[inwin], f[inwin], lw=2.45, color=trace_color)
        if np.any(peaks):
            ax.scatter(
                x[peaks], f[peaks], marker="^", s=96,
                facecolors=QC_COLORS["peak"], edgecolors="black", linewidths=0.85,
                zorder=6,
            )
        if np.any(troughs):
            ax.scatter(
                x[troughs], f[troughs], marker="v", s=96,
                facecolors=QC_COLORS["trough"], edgecolors="black", linewidths=0.85,
                zorder=6,
            )
        xmin, xmax = float(np.nanmin(x[inwin])), float(np.nanmax(x[inwin]))
        pad = max(1.0, 0.08 * (xmax - xmin))
        ax.set_xlim(xmin - pad, xmax + pad)
    else:
        ax.plot(x, f, lw=2.0, color=trace_color)

    for key, label, style, color in [
        ("fc_N", "Fc", "-.", QC_COLORS["fc"]),
        ("fci_N", "Fci", ":", QC_COLORS["fci"]),
        ("fmean_N", "Fmean", "--", QC_COLORS["fmean"]),
    ]:
        val = result.get(key, np.nan)
        if np.isfinite(val):
            ax.axhline(val, ls=style, lw=1.9, color=color)

    if review_required:
        txt = (
            f"Review: {quality_flag}\n"
            f"Peak span={result.get('peak_spread_mm', np.nan):.2f} mm; "
            f"trough span={result.get('trough_spread_mm', np.nan):.2f} mm"
        )
        ax.text(
            0.025, 0.955, txt, transform=ax.transAxes, fontsize=8.4, ha="left", va="top",
            bbox=dict(boxstyle="round,pad=0.32", facecolor="#fff8e1", edgecolor=QC_COLORS["drop"], linewidth=1.05, alpha=0.99),
        )
    ax.set_title("Zoomed selected window + extraction points", fontsize=12, fontweight="bold", pad=8)
    ax.set_xlabel("Displacement (mm)")
    ax.set_ylabel("Force (N)")
    ax.grid(True, alpha=0.15)

    # Shared figure-level legend: outside plotting area to avoid covering traces.
    legend_handles = [
        Line2D([0], [0], color=trace_color, lw=2.0, label="Force trace"),
        Patch(facecolor=QC_COLORS["selected_window"], edgecolor="none", alpha=0.55, label="Selected window"),
        Line2D([0], [0], color=QC_COLORS["window_edge"], lw=1.45, ls="--", label="Window start/end"),
        Line2D([0], [0], color=QC_COLORS["drop"], lw=2.0, ls=":", label="Break/drop proxy"),
        Line2D([0], [0], color=QC_COLORS["fc"], lw=1.9, ls="-.", label=f"Fc={result.get('fc_N', np.nan):.3g} N" if np.isfinite(result.get('fc_N', np.nan)) else "Fc"),
        Line2D([0], [0], color=QC_COLORS["fci"], lw=1.9, ls=":", label=f"Fci={result.get('fci_N', np.nan):.3g} N" if np.isfinite(result.get('fci_N', np.nan)) else "Fci"),
        Line2D([0], [0], color=QC_COLORS["fmean"], lw=1.9, ls="--", label=f"Fmean={result.get('fmean_N', np.nan):.3g} N" if np.isfinite(result.get('fmean_N', np.nan)) else "Fmean"),
        Line2D([0], [0], marker="^", color="black", markerfacecolor=QC_COLORS["peak"], lw=0, markersize=8, label="Top5 peaks"),
        Line2D([0], [0], marker="v", color="black", markerfacecolor=QC_COLORS["trough"], lw=0, markersize=8, label="Bot5 troughs"),
    ]
    fig.legend(
        handles=legend_handles,
        loc="lower center",
        bbox_to_anchor=(0.5, 0.035),
        ncol=5,
        frameon=True,
        framealpha=0.98,
        facecolor="white",
        edgecolor="0.65",
        fontsize=8.0,
    )

    fig.savefig(outpath, dpi=config.dpi, facecolor="white", edgecolor="white", transparent=False, bbox_inches="tight", pad_inches=0.12)
    plt.close(fig)


# =============================================================================
# Scanning, mapping, summaries
# =============================================================================

def read_full_table_for_scan(path: str, sheet: str) -> pd.DataFrame:
    """Read a complete table for scan/preflight.

    The manuscript-guided notebook uses this to estimate the available displacement span
    before asking the user to proceed. It remains intentionally simple: one header row,
    then numeric data exported by a mechanical tester.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext in {".xlsx", ".xlsm", ".xls"}:
        return read_table_from_excel_sheet(path, sheet, header_row=1)
    return pd.read_csv(path)


def scan_workbook(config: ManuscriptConfig) -> Dict[str, Any]:
    """Scan an input workbook before running manuscript metrics.

    This scan is intentionally conservative: it checks force/displacement column plausibility,
    unit consistency, user geometry, and whether the fixed 25.0 mm manuscript window is likely
    feasible. It does not calculate final metrics.
    """
    outdir = ensure_dir(os.path.join(config.output_root, f"scan_{now_stamp()}"))
    sheets = workbook_sheet_names(config.input_path)
    rows = []
    final_audit_all = []
    audit_all = []
    min_span_needed_mm = (config.start_offset_mm + config.window_length_mm) / max(1e-12, (1.0 - config.tail_exclude_frac))
    for sheet in sheets:
        include = not any(term.lower() in sheet.lower() for term in config.exclude_sheet_if_contains)
        try:
            full_df = read_full_table_for_scan(config.input_path, sheet)
            try:
                _WORKBOOK_DATA_CACHE[(os.path.abspath(config.input_path), sheet)] = full_df
            except Exception:
                pass
            preview = full_df.head(500)
            assignment = detect_force_displacement_columns(preview)
            force_unit, disp_unit, unit_warnings, contradiction = unit_resolution(config, assignment)
            group, adhesive, rep = infer_group_and_adhesive(sheet)
            nrows = len(full_df)
            # Full-span rough displacement range in mm, using the detected displacement column.
            xspan_header_unit = np.nan
            xspan_mm = np.nan
            x_min_mm = np.nan
            x_max_mm = np.nan
            if assignment.displacement_column and assignment.displacement_column in full_df.columns:
                arr = numeric_array(full_df[assignment.displacement_column])
                arr = arr[np.isfinite(arr)]
                if arr.size > 1:
                    raw_min = float(np.nanmin(arr))
                    raw_max = float(np.nanmax(arr))
                    xspan_header_unit = raw_max - raw_min
                    if disp_unit:
                        converted = unit_convert_displacement(np.array([raw_min, raw_max], dtype=float), disp_unit)
                        x_min_mm = float(np.nanmin(converted))
                        x_max_mm = float(np.nanmax(converted))
                        xspan_mm = float(x_max_mm - x_min_mm)
            problem_parts = []
            warning_parts = list(unit_warnings)
            if assignment.needs_user_confirmation:
                problem_parts.append("Column detection confidence is LOW; user confirmation required.")
            if contradiction and config.stop_on_unit_contradiction:
                problem_parts.append("Unit contradiction detected.")
            if not force_unit or not disp_unit:
                problem_parts.append("Unit unresolved.")
            if config.width_mm is None or float(config.width_mm) <= 0:
                problem_parts.append("Specimen width_mm must be confirmed and > 0.")
            if config.thickness_mm is None or float(config.thickness_mm) <= 0:
                problem_parts.append("Specimen thickness_mm must be confirmed and > 0.")
            if config.bond_length_mm is None or float(config.bond_length_mm) <= 0:
                problem_parts.append("Bond length/overlap must be confirmed and > 0.")
            if np.isfinite(xspan_mm) and xspan_mm < min_span_needed_mm:
                warning_parts.append(
                    f"The fixed manuscript window may be infeasible: displacement span {xspan_mm:.2f} mm < approximate needed span {min_span_needed_mm:.2f} mm."
                )
            if config.bond_length_mm and config.bond_length_mm < config.window_length_mm:
                warning_parts.append(
                    f"Bond length ({config.bond_length_mm:.2f} mm) is shorter than the manuscript window ({config.window_length_mm:.2f} mm); verify whether this is intentional."
                )
            rows.append({
                "sheet": sheet,
                "include": include,
                "main_group": group,
                "adhesive": adhesive,
                "replicate_id": rep,
                "force_column": assignment.force_column,
                "displacement_column": assignment.displacement_column,
                "column_confidence": assignment.confidence,
                "column_pair_score": assignment.pair_score,
                "column_reason": assignment.reason,
                "force_unit": force_unit,
                "displacement_unit": disp_unit,
                "width_mm": config.width_mm,
                "thickness_mm": config.thickness_mm,
                "bond_length_mm": config.bond_length_mm,
                "notes": "",
                "rows_detected": nrows,
                "rough_displacement_span_header_unit": xspan_header_unit,
                "rough_displacement_min_mm": x_min_mm,
                "rough_displacement_max_mm": x_max_mm,
                "rough_displacement_span_mm": xspan_mm,
                "minimum_span_needed_for_25mm_window_mm": min_span_needed_mm,
                "warnings": "; ".join(warning_parts),
                "problem": " ".join(problem_parts),
            })
            final_audit_all.append({
                "sheet": sheet,
                "selected_force_column": assignment.force_column,
                "selected_displacement_column": assignment.displacement_column,
                "force_unit": force_unit,
                "displacement_unit": disp_unit,
                "confidence": assignment.confidence,
                "pair_score": round(float(assignment.pair_score), 3),
                "needs_user_confirmation": bool(assignment.needs_user_confirmation),
                "decision": "ACCEPTED" if not assignment.needs_user_confirmation and not problem_parts else "REVIEW_REQUIRED",
                "reason": assignment.reason,
                "warnings": "; ".join(warning_parts),
                "problem": " ".join(problem_parts),
            })
            for ar in assignment.audit_rows:
                ar2 = dict(ar)
                ar2["sheet"] = sheet
                audit_all.append(ar2)
            del full_df
            gc.collect()
        except Exception as exc:
            rows.append({"sheet": sheet, "include": include, "problem": f"Scan failed: {exc}"})
    mapping = pd.DataFrame(rows)
    final_audit = pd.DataFrame(final_audit_all)
    audit = pd.DataFrame(audit_all)
    mapping_file = os.path.join(outdir, "sheet_mapping_review.csv")
    audit_file = os.path.join(outdir, "column_assignment_audit.csv")
    advanced_audit_file = os.path.join(outdir, "advanced_column_pair_scores.csv")
    mapping.to_csv(mapping_file, index=False)
    final_audit.to_csv(audit_file, index=False)
    audit.to_csv(advanced_audit_file, index=False)

    critical = mapping["problem"].fillna("").astype(str).str.len() > 0 if not mapping.empty and "problem" in mapping else []
    warnings_mask = mapping["warnings"].fillna("").astype(str).str.len() > 0 if not mapping.empty and "warnings" in mapping else []
    n_critical = int(np.sum(critical)) if len(mapping) else 0
    n_warning = int(np.sum(warnings_mask)) if len(mapping) else 0
    n_included = int(mapping.get("include", pd.Series(dtype=bool)).astype(bool).sum()) if not mapping.empty and "include" in mapping else 0
    score = max(0, 100 - 15 * n_critical - 3 * n_warning)
    min_span = pd.to_numeric(mapping.get("rough_displacement_span_mm", pd.Series(dtype=float)), errors="coerce").min() if not mapping.empty else np.nan
    max_span = pd.to_numeric(mapping.get("rough_displacement_span_mm", pd.Series(dtype=float)), errors="coerce").max() if not mapping.empty else np.nan
    readiness = pd.DataFrame([{
        "file": os.path.basename(config.input_path),
        "sheets_detected": len(sheets),
        "sheets_included": n_included,
        "critical_problem_count": n_critical,
        "warning_count": n_warning,
        "readiness_score_0_100": score,
        "min_detected_displacement_span_mm": min_span,
        "max_detected_displacement_span_mm": max_span,
        "approx_minimum_span_needed_for_25mm_window_mm": min_span_needed_mm,
        "recommended_action": "Proceed after user confirmation" if n_critical == 0 else "Review/edit sheet_mapping_review.csv before analysis",
        "estimated_runtime_note": "Manuscript mode is usually fast; QC-plot export increases runtime for large workbooks.",
    }])
    readiness_file = os.path.join(outdir, "readiness_report.csv")
    readiness.to_csv(readiness_file, index=False)
    return {
        "outdir": outdir,
        "sheet_mapping_file": mapping_file,
        "column_assignment_audit_file": audit_file,
        "advanced_column_pair_scores_file": advanced_audit_file,
        "readiness_report_file": readiness_file,
        "mapping": mapping,
        "audit": final_audit,
        "advanced_audit": audit,
        "readiness": readiness,
    }


def load_sheet_mapping(config: ManuscriptConfig) -> pd.DataFrame:
    if config.sheet_mapping_file and os.path.exists(config.sheet_mapping_file):
        return pd.read_csv(config.sheet_mapping_file)
    scan = scan_workbook(config)
    return scan["mapping"]


def build_trace_inputs(config: ManuscriptConfig, mapping: pd.DataFrame) -> List[TraceInput]:
    """Build TraceInput objects from a confirmed mapping table."""
    traces: List[TraceInput] = []
    ext = os.path.splitext(config.input_path)[1].lower()
    for _, row in mapping.iterrows():
        sheet = str(row.get("sheet", ""))
        include = parse_bool(row.get("include", True), default=True)
        if not include:
            continue
        force_col = str(row.get("force_column", ""))
        disp_col = str(row.get("displacement_column", ""))
        if not force_col or not disp_col:
            continue
        if ext in {".xlsx", ".xlsm", ".xls"}:
            cache_key = (os.path.abspath(config.input_path), sheet)
            if cache_key in _WORKBOOK_DATA_CACHE:
                df = _WORKBOOK_DATA_CACHE[cache_key].copy(deep=False)
            else:
                df = read_table_from_excel_sheet(config.input_path, sheet, header_row=1)
        else:
            df = pd.read_csv(config.input_path)
        traces.append(TraceInput(
            sheet=sheet,
            main_group=str(row.get("main_group", "UNKNOWN")),
            adhesive=str(row.get("adhesive", "UNKNOWN")),
            replicate_id=str(row.get("replicate_id", sheet)),
            include=include,
            width_mm=float(row.get("width_mm", config.width_mm) or config.width_mm),
            thickness_mm=float(row.get("thickness_mm", config.thickness_mm) or config.thickness_mm),
            bond_length_mm=float(row.get("bond_length_mm", config.bond_length_mm) or config.bond_length_mm),
            force_column=force_col,
            displacement_column=disp_col,
            force_unit=str(row.get("force_unit", config.force_unit) or config.force_unit),
            displacement_unit=str(row.get("displacement_unit", config.displacement_unit) or config.displacement_unit),
            data=df,
            column_confidence=str(row.get("column_confidence", "")),
            column_reason=str(row.get("column_reason", "")),
            warnings=[str(row.get("warnings", ""))] if str(row.get("warnings", "")) not in {"", "nan"} else [],
        ))
        gc.collect()
    return traces


def summarize_groups(results: pd.DataFrame) -> pd.DataFrame:
    if results.empty:
        return pd.DataFrame()
    pass_df = results[results["status"] == "PASS"].copy()
    if pass_df.empty:
        return pd.DataFrame()
    rows = []
    group_cols = ["main_group", "adhesive"]
    for keys, g in pass_df.groupby(group_cols, dropna=False):
        d = {"main_group": keys[0], "adhesive": keys[1], "n_pass": len(g)}
        for col in ["fc_over_w_N_per_mm", "fci_over_w_N_per_mm", "fc_N", "fci_N", "ssa_abs_N", "ssa_norm", "psi_mu_over_sigma", "window_cv_pct", "displacement_at_break_mm"]:
            vals = pd.to_numeric(g[col], errors="coerce").dropna()
            d[f"{col}_mean"] = vals.mean() if len(vals) else np.nan
            d[f"{col}_sd"] = vals.std(ddof=1) if len(vals) > 1 else np.nan
            d[f"{col}_cv_pct"] = 100 * vals.std(ddof=1) / vals.mean() if len(vals) > 1 and vals.mean() != 0 else np.nan
        # Direct CoV_Fci used in the paper.
        vals = pd.to_numeric(g["fci_N"], errors="coerce").dropna()
        d["CoV_Fci_pct"] = 100 * vals.std(ddof=1) / vals.mean() if len(vals) > 1 and vals.mean() != 0 else np.nan
        rows.append(d)
    return pd.DataFrame(rows)


def formula_dictionary() -> pd.DataFrame:
    return pd.DataFrame([
        {"metric": "Fc", "formula": "mean(top 5 local maxima and bottom 5 local minima)", "latex": r"F_c = \frac{\sum_{i=1}^{5}F_{p,i}+\sum_{j=1}^{5}F_{v,j}}{10}", "unit": "N", "interpretation": "Representative peel force from the selected manuscript window."},
        {"metric": "Fc/w", "formula": "Fc / w", "latex": r"F_c/w", "unit": "N mm^-1", "interpretation": "Width-normalized peel/bond strength used for conventional comparison."},
        {"metric": "Fci", "formula": "mean(top 5 local maxima)", "latex": r"F_{ci}=\frac{1}{5}\sum_{i=1}^{5}F_{p,i}", "unit": "N", "interpretation": "Crack-initiation force descriptor for oscillatory or restart-like traces."},
        {"metric": "Fci/w", "formula": "Fci / w", "latex": r"F_{ci}/w", "unit": "N mm^-1", "interpretation": "Width-normalized initiation-force descriptor."},
        {"metric": "PSI", "formula": "mean(F_window) / sample_SD(F_window)", "latex": r"PSI=\bar{F}_{win}/s_{F,win}", "unit": "dimensionless", "interpretation": "Higher values indicate steadier force response within the selected window."},
        {"metric": "SSA", "formula": "mean(top 5 peaks) - mean(bottom 5 troughs)", "latex": r"SSA=\bar{F}_p-\bar{F}_v", "unit": "N", "interpretation": "Operational force-oscillation amplitude descriptor; not intrinsic fracture energy."},
        {"metric": "SSA_norm", "formula": "SSA / mean(F_window)", "latex": r"SSA_{norm}=SSA/\bar{F}_{win}", "unit": "dimensionless", "interpretation": "Force-oscillation amplitude normalized by force scale."},
        {"metric": "window_CV", "formula": "100 * sample_SD(F_window) / mean(F_window)", "latex": r"CV_{win}=100s_{F,win}/\bar{F}_{win}", "unit": "%", "interpretation": "Relative force fluctuation in selected window."},
        {"metric": "drift_ratio", "formula": "|m| L_win / Fbar_win", "latex": r"|m|L_{win}/\bar{F}_{win}\le 0.25", "unit": "dimensionless", "interpretation": "Candidate-window ramping filter; manuscript limit <= 0.25."},
        {"metric": "sigma_ci", "formula": "Fci / (w t)", "latex": r"\sigma_{ci}=F_{ci}/(wt)", "unit": "MPa when N/mm^2", "interpretation": "Optional nominal peel-arm axial stress screen; not a full stress analysis."},
    ])


def metric_guide() -> pd.DataFrame:
    return pd.DataFrame([
        {"topic": "Purpose", "description": "This notebook reproduces the manuscript force-trace workflow for soft textile T-peel screening."},
        {"topic": "System-dependence caution", "description": "Peel metrics are protocol-defined descriptors. They may include effects of geometry, peel-arm deformation, fixture compliance, and force-trace instability."},
        {"topic": "SSA caution", "description": "SSA compares peak-valley force oscillation under the same processing protocol. It should not be treated as an intrinsic crack-front property without additional evidence."},
        {"topic": "Strict extrema rule", "description": "In manuscript_baseline_v1 mode, Top5/Bottom5-derived metrics are reported only when >=5 peaks and >=5 troughs are available in the selected window."},
        {"topic": "Recommendation mode", "description": "Recommendation scans are not part of default manuscript replication. They are only a user-support/recovery layer if baseline analysis fails."},
    ])


def output_file_guide() -> pd.DataFrame:
    """Describe the default output files for non-specialist users."""
    return pd.DataFrame([
        {"output": "manuscript_baseline_outputs.zip", "format": "ZIP", "contents": "All core manuscript-baseline/user-profile outputs from the run.", "use": "Single archive for sharing, storage, and reproducibility."},
        {"output": "analysis_report.xlsx", "format": "Excel", "contents": "Readiness, paper metrics, group summary, formulas, metric guide, validation, and input mapping.", "use": "Main human-readable report."},
        {"output": "paper_metrics.csv", "format": "CSV", "contents": "Per-trace manuscript metrics: Fc, Fci, Fc/w, Fci/w, PSI, SSA, selected window, peak/trough counts, and warnings.", "use": "Main machine-readable result table."},
        {"output": "group_summary.csv", "format": "CSV", "contents": "Mean, sample SD, CoV, and n by group/adhesive for passing traces.", "use": "Manuscript-style replicate summary."},
        {"output": "input_audit.csv", "format": "CSV", "contents": "Detected sheets, inclusion flags, columns, units, geometry values, displacement span checks, and warnings.", "use": "Traceability and troubleshooting."},
        {"output": "column_assignment_audit.csv", "format": "CSV", "contents": "Final selected force/displacement assignment per sheet with confidence and reason.", "use": "User-facing confirmation of the columns actually used."},
        {"output": "advanced_column_pair_scores.csv", "format": "CSV", "contents": "All candidate force/displacement column-pair scores tested during detection.", "use": "Advanced debugging only; not needed for routine interpretation."},
        {"output": "output_consistency_audit.csv", "format": "CSV", "contents": "Checks that CSV and Excel exports match the in-memory calculated result tables.", "use": "Credibility and data-integrity guardrail."},
        {"output": "diagnostic_summary.csv", "format": "CSV", "contents": "Separates analysis status, metric-quality review status, export integrity, and built-in validation status.", "use": "Start here when deciding whether outputs are safe to interpret."},
        {"output": "issue_action_table.csv", "format": "CSV", "contents": "Maps common failures/warnings to conservative next actions.", "use": "Guidance for failed traces or caution flags."},
        {"output": "trace_integrity_audit.csv", "format": "CSV", "contents": "Checks missing values, monotonic displacement behavior, duplicate displacement points, and large displacement gaps.", "use": "Preflight data-quality audit before interpreting trace metrics."},
        {"output": "run_history.csv", "format": "CSV", "contents": "Run type, profile hash, input hash, elapsed time, pass/fail/review counts, and output folder.", "use": "Documents reruns and sensitivity/profile comparisons."},
        {"output": "validation_count_summary.csv", "format": "CSV", "contents": "Explains the count and purpose of validation/audit checks.", "use": "Shows what the PASS counts actually mean."},
        {"output": "formula_dictionary.csv", "format": "CSV", "contents": "Metric equations, units, and interpretations.", "use": "Prevents ambiguity in column meanings."},
        {"output": "metric_guide.csv", "format": "CSV", "contents": "Plain-language interpretation notes and limitations.", "use": "Helps non-specialist users interpret outputs."},
        {"output": "scotch_validation.csv", "format": "CSV", "contents": "Built-in Scotch Tape T-peel validation checks when the included dataset is run.", "use": "Confirms the notebook environment reproduces expected benchmark outputs."},
        {"output": "qc_plots/", "format": "PNG folder", "contents": "Full-trace and zoomed selected-window plots with extraction points.", "use": "Visual quality control for selected windows and Top5/Bottom5 extraction."},
        {"output": "provenance.json", "format": "JSON", "contents": "Software version, timestamp, input file hash, and locked settings.", "use": "Reproducibility record."},
        {"output": "user_method_profile.json", "format": "JSON", "contents": "Optional user-created analysis profile for NON_MANUSCRIPT_MODIFIED runs.", "use": "Reusable profile for future datasets; does not overwrite manuscript_baseline_v1."},
        {"output": "user_benchmark_expected.csv", "format": "CSV", "contents": "Optional user-created expected-value table from a trusted run.", "use": "Validate future reruns of the same dataset/profile without changing the built-in Scotch reference."},
        {"output": "user_benchmark_metadata.json", "format": "JSON", "contents": "Optional metadata for user-created benchmark files.", "use": "Documents benchmark origin, software version, selected sheets, and metrics."},
    ])



def software_versions() -> Dict[str, str]:
    """Return core package versions for provenance."""
    import sys as _sys
    versions = {
        "python": _sys.version.split()[0],
        "pandas": getattr(pd, "__version__", ""),
        "numpy": getattr(np, "__version__", ""),
        "matplotlib": getattr(plt.matplotlib, "__version__", ""),
    }
    try:
        import scipy as _scipy
        versions["scipy"] = getattr(_scipy, "__version__", "")
    except Exception:
        versions["scipy"] = "not available"
    try:
        import openpyxl as _openpyxl
        versions["openpyxl"] = getattr(_openpyxl, "__version__", "")
    except Exception:
        versions["openpyxl"] = "not available"
    return versions


# =============================================================================
# Scotch validation
# =============================================================================

SCOTCH_EXPECTED = pd.DataFrame([
    {"sheet": "Scotch Tape T Peel_Jul8_Test2", "fc_N": 4.036280679, "fci_N": 4.366416781, "fc_over_w_N_per_mm": 0.2124358252, "fci_over_w_N_per_mm": 0.2298114096, "window_cv_pct": 5.197600715, "ssa_abs_N": 0.6602722042, "psi_mu_over_sigma": 19.23964642, "window_start_mm": 20.46, "window_end_mm": 45.46, "peak_count": 27, "trough_count": 27},
    {"sheet": "Scotch Tape T Peel_Jul8_Test3", "fc_N": 3.480467086, "fci_N": 3.952945378, "fc_over_w_N_per_mm": 0.1831824782, "fci_over_w_N_per_mm": 0.2080497567, "window_cv_pct": 7.745181595, "ssa_abs_N": 0.9449565851, "psi_mu_over_sigma": 12.91125312, "window_start_mm": 18.653, "window_end_mm": 43.653, "peak_count": 31, "trough_count": 31},
    {"sheet": "Scotch Tape T Peel_Jul8_Test4", "fc_N": 3.115101073, "fci_N": 3.466622911, "fc_over_w_N_per_mm": 0.1639526880, "fci_over_w_N_per_mm": 0.1824538374, "window_cv_pct": 7.249375627, "ssa_abs_N": 0.7030436767, "psi_mu_over_sigma": 13.79429142, "window_start_mm": 7.886, "window_end_mm": 32.886, "peak_count": 34, "trough_count": 34},
])


def validate_scotch(results: pd.DataFrame, rel_tol: float = 0.015, abs_tol: float = 0.015) -> pd.DataFrame:
    rows = []
    for _, exp in SCOTCH_EXPECTED.iterrows():
        sheet = exp["sheet"]
        actual_rows = results[results["sheet"] == sheet]
        if actual_rows.empty:
            rows.append({"sheet": sheet, "validation_status": "FAIL", "metric": "sheet_present", "expected": "present", "actual": "missing", "difference": np.nan})
            continue
        act = actual_rows.iloc[0]
        for col in ["fc_N", "fci_N", "fc_over_w_N_per_mm", "fci_over_w_N_per_mm", "window_cv_pct", "ssa_abs_N", "psi_mu_over_sigma", "window_start_mm", "window_end_mm", "peak_count", "trough_count"]:
            e = exp[col]
            a = act.get(col, np.nan)
            try:
                diff = float(a) - float(e)
                ok = abs(diff) <= max(abs_tol, rel_tol * max(abs(float(e)), 1e-12))
            except Exception:
                diff = np.nan
                ok = False
            rows.append({"sheet": sheet, "metric": col, "expected": e, "actual": a, "difference": diff, "validation_status": "PASS" if ok else "FAIL"})
    return pd.DataFrame(rows)


# =============================================================================
# Output consistency audit
# =============================================================================

def _compare_dataframe_schema_and_values(name: str, expected: pd.DataFrame, observed: pd.DataFrame, atol: float = 1e-9, rtol: float = 1e-9) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    rows.append({"check": f"{name}: row count", "status": "PASS" if len(expected) == len(observed) else "FAIL", "details": f"expected {len(expected)}, observed {len(observed)}"})
    missing_cols = [c for c in expected.columns if c not in observed.columns]
    extra_cols = [c for c in observed.columns if c not in expected.columns]
    rows.append({"check": f"{name}: column set", "status": "PASS" if not missing_cols else "FAIL", "details": f"missing={missing_cols}; extra={extra_cols}"})
    common_cols = [c for c in expected.columns if c in observed.columns]
    if len(expected) != len(observed):
        return rows
    for col in common_cols:
        e = expected[col].reset_index(drop=True)
        o = observed[col].reset_index(drop=True)
        # Treat boolean columns as categorical/text for audit purposes.
        if str(e.dtype) == "bool" or str(o.dtype) == "bool" or set(e.dropna().unique()).issubset({True, False}) or set(o.dropna().unique()).issubset({True, False}):
            ok = e.fillna("").astype(str).equals(o.fillna("").astype(str))
            rows.append({"check": f"{name}: boolean/text column {col}", "status": "PASS" if ok else "FAIL", "details": "values match" if ok else "values differ"})
            continue
        e_num = pd.to_numeric(e, errors="coerce")
        o_num = pd.to_numeric(o, errors="coerce")
        numeric_like = e_num.notna().sum() > 0 or o_num.notna().sum() > 0
        if numeric_like:
            mask = e_num.notna() | o_num.notna()
            if mask.sum() == 0:
                ok = True
                maxdiff = 0.0
            else:
                diff = (e_num[mask] - o_num[mask]).abs()
                denom = np.maximum(e_num[mask].abs(), 1e-12)
                ok = bool(((diff <= (atol + rtol * denom)) | (e_num[mask].isna() & o_num[mask].isna())).all())
                maxdiff = float(diff.max()) if len(diff) else 0.0
            rows.append({"check": f"{name}: numeric column {col}", "status": "PASS" if ok else "FAIL", "details": f"max_abs_diff={maxdiff:.3g}"})
        else:
            ok = e.fillna("").astype(str).equals(o.fillna("").astype(str))
            rows.append({"check": f"{name}: text column {col}", "status": "PASS" if ok else "FAIL", "details": "text values match" if ok else "text values differ"})
    return rows


def run_output_consistency_audit(
    results_df: pd.DataFrame,
    group_df: pd.DataFrame,
    paper_metrics_file: str,
    group_file: str,
    excel_file: str = "",
) -> pd.DataFrame:
    """Verify that CSV and Excel exports match the in-memory analysis tables.

    This audit is deliberately conservative. It catches accidental data-writing mismatches,
    dropped columns, row-count changes, and silent numeric changes between the core result
    tables and exported files.
    """
    rows: List[Dict[str, Any]] = []
    try:
        observed_results = pd.read_csv(paper_metrics_file)
        rows.extend(_compare_dataframe_schema_and_values("paper_metrics.csv", results_df, observed_results))
    except Exception as exc:
        rows.append({"check": "paper_metrics.csv readable", "status": "FAIL", "details": str(exc)})
    try:
        observed_group = pd.read_csv(group_file)
        rows.extend(_compare_dataframe_schema_and_values("group_summary.csv", group_df, observed_group))
    except Exception as exc:
        rows.append({"check": "group_summary.csv readable", "status": "FAIL", "details": str(exc)})
    if excel_file and os.path.exists(excel_file):
        try:
            excel_results = pd.read_excel(excel_file, sheet_name="Paper_Metrics")
            rows.extend(_compare_dataframe_schema_and_values("Excel Paper_Metrics", results_df, excel_results))
        except Exception as exc:
            rows.append({"check": "Excel Paper_Metrics readable", "status": "FAIL", "details": str(exc)})
        try:
            excel_group = pd.read_excel(excel_file, sheet_name="Group_Summary")
            rows.extend(_compare_dataframe_schema_and_values("Excel Group_Summary", group_df, excel_group))
        except Exception as exc:
            rows.append({"check": "Excel Group_Summary readable", "status": "FAIL", "details": str(exc)})
    return pd.DataFrame(rows)



def trace_integrity_audit_from_traces(traces: Sequence[TraceInput], config: ManuscriptConfig) -> pd.DataFrame:
    """Preflight force-displacement trace integrity audit.

    The audit is intentionally tolerant of normal nonuniform sampling from mechanical testers.
    It flags severe issues that would make windowing, extrema detection, or displacement-at-break
    interpretation unreliable: missing force/displacement values, nearly constant displacement,
    duplicate displacement points, large gaps, and non-monotonic displacement progression.
    """
    rows: List[Dict[str, Any]] = []
    for trace in traces:
        try:
            x_raw = pd.to_numeric(trace.data[trace.displacement_column], errors="coerce").to_numpy(dtype=float)
            f_raw = pd.to_numeric(trace.data[trace.force_column], errors="coerce").to_numpy(dtype=float)
            n_total = int(max(len(x_raw), len(f_raw)))
            valid = np.isfinite(x_raw) & np.isfinite(f_raw)
            x = x_raw[valid]
            f = f_raw[valid]
            x = unit_convert_displacement(x, trace.displacement_unit)
            f = unit_convert_force(f, trace.force_unit)
            if len(x) >= 2 and (np.nanmax(x) - np.nanmin(x)) < 0 and abs(np.nanmin(x)) > abs(np.nanmax(x)):
                x = -x
            # Use displacement differences after unit conversion and sign correction.
            dx = np.diff(x) if len(x) >= 2 else np.array([])
            abs_dx = np.abs(dx[np.isfinite(dx)])
            positive = dx > 0
            negative = dx < 0
            zero = dx == 0
            median_step = float(np.nanmedian(abs_dx[abs_dx > 0])) if np.any(abs_dx > 0) else np.nan
            large_gap_count = int(np.sum(abs_dx > 10.0 * median_step)) if np.isfinite(median_step) and median_step > 0 else 0
            max_gap_ratio = float(np.nanmax(abs_dx) / median_step) if np.isfinite(median_step) and median_step > 0 and len(abs_dx) else np.nan
            monotonic_frac = float(max(np.sum(positive | zero), np.sum(negative | zero)) / len(dx)) if len(dx) else np.nan
            duplicate_frac = float(np.sum(zero) / len(dx)) if len(dx) else np.nan
            usable_span = float(np.nanmax(x) - np.nanmin(x)) if len(x) else np.nan
            n_missing_force = int(np.sum(~np.isfinite(f_raw)))
            n_missing_disp = int(np.sum(~np.isfinite(x_raw)))
            messages = []
            status = "PASS"
            if len(x) < max(10, config.min_window_points):
                status = "FAIL"; messages.append("Too few valid force-displacement points for analysis.")
            if not np.isfinite(usable_span) or usable_span <= 0:
                status = "FAIL"; messages.append("No usable displacement span detected.")
            if np.isfinite(monotonic_frac) and monotonic_frac < 0.90:
                status = "FAIL" if monotonic_frac < 0.70 else "WARN"; messages.append("Displacement is not mostly monotonic; check skipped/reordered data.")
            if np.isfinite(duplicate_frac) and duplicate_frac > 0.20:
                status = "FAIL" if duplicate_frac > 0.50 else "WARN"; messages.append("High fraction of repeated displacement values.")
            if large_gap_count > 0:
                if status == "PASS": status = "WARN"
                messages.append(f"Detected {large_gap_count} displacement gap(s) >10× median step; possible skipped data/logging discontinuity.")
            if n_missing_force or n_missing_disp:
                if status == "PASS": status = "WARN"
                messages.append("Missing force/displacement values were removed before analysis.")
            rows.append({
                "sheet": trace.sheet,
                "n_points_raw": n_total,
                "n_valid_points": int(len(x)),
                "n_missing_force": n_missing_force,
                "n_missing_displacement": n_missing_disp,
                "duplicate_displacement_fraction": duplicate_frac,
                "monotonic_displacement_fraction": monotonic_frac,
                "usable_displacement_span_mm": usable_span,
                "median_displacement_step_mm": median_step,
                "max_gap_to_median_step_ratio": max_gap_ratio,
                "large_gap_count": large_gap_count,
                "trace_integrity_status": status,
                "trace_integrity_message": "; ".join(messages) if messages else "Trace passed continuity/missing-data preflight checks.",
            })
        except Exception as exc:
            rows.append({
                "sheet": getattr(trace, "sheet", ""),
                "n_points_raw": np.nan,
                "n_valid_points": np.nan,
                "n_missing_force": np.nan,
                "n_missing_displacement": np.nan,
                "duplicate_displacement_fraction": np.nan,
                "monotonic_displacement_fraction": np.nan,
                "usable_displacement_span_mm": np.nan,
                "median_displacement_step_mm": np.nan,
                "max_gap_to_median_step_ratio": np.nan,
                "large_gap_count": np.nan,
                "trace_integrity_status": "FAIL",
                "trace_integrity_message": f"Trace-integrity audit failed: {exc}",
            })
    return pd.DataFrame(rows)


def require_no_failed_trace_integrity(trace_integrity_df: pd.DataFrame) -> None:
    """Raise if trace-integrity audit has hard failures.

    The public notebook may choose to continue with non-failing sheets, but this helper is
    available when a strict fail-fast workflow is preferred.
    """
    if trace_integrity_df is not None and not trace_integrity_df.empty:
        n_fail = int((trace_integrity_df.get("trace_integrity_status", pd.Series(dtype=str)) == "FAIL").sum())
        if n_fail:
            raise RuntimeError(f"Trace-integrity preflight found {n_fail} failed sheet(s). Review trace_integrity_audit.csv before analysis.")

# =============================================================================
# Full pipeline
# =============================================================================

def run_manuscript_pipeline(config: ManuscriptConfig, progress_callback: Optional[Any] = None) -> Dict[str, Any]:
    baseline_check = validate_against_locked_manuscript_profile(config)
    if baseline_check["is_manuscript_baseline"]:
        config.method_profile = "manuscript_baseline_v1"
    else:
        if not getattr(config, "allow_non_manuscript_run", False):
            details = "; ".join(
                f"{row['parameter']}: expected {row['expected']}, found {row['actual']}"
                for row in baseline_check["changed_parameters"]
            )
            raise ValueError(
                "Non-manuscript settings detected. Use build_manuscript_baseline_config() "
                "or set allow_non_manuscript_run=True for an explicitly labeled exploratory run. "
                + details
            )
        config.method_profile = "non_manuscript_modified"

    outdir = ensure_dir(os.path.join(config.output_root, f"manuscript_run_{now_stamp()}"))
    plots_dir = ensure_dir(os.path.join(outdir, "qc_plots"))

    if not config.sheet_mapping_file:
        scan = scan_workbook(config)
        mapping = scan["mapping"]
        # Keep scan files accessible in run folder too.
        mapping_file = scan["sheet_mapping_file"]
        column_audit_file = scan["column_assignment_audit_file"]
        readiness_file = scan["readiness_report_file"]
    else:
        mapping = load_sheet_mapping(config)
        mapping_file = config.sheet_mapping_file
        column_audit_file = ""
        readiness_file = ""

    traces = build_trace_inputs(config, mapping)
    trace_integrity_df = trace_integrity_audit_from_traces(traces, config)
    results: List[Dict[str, Any]] = []
    trace_tables: Dict[str, pd.DataFrame] = {}
    start_time = time.time()
    for idx, trace in enumerate(traces, start=1):
        message = f"[{idx}/{len(traces)}] Analyzing {trace.sheet} ..."
        if progress_callback is not None:
            try:
                progress_callback(idx, len(traces), trace.sheet, "analyzing")
            except Exception:
                pass
        else:
            print(message)
        res, trace_df = analyze_trace(trace, config)
        results.append(res)
        if trace_df is not None:
            trace_tables[trace.sheet] = trace_df
            if config.export_qc_plots:
                plot_name = f"{safe_filename(trace.sheet)}_qc.png"
                make_qc_plot(res, trace_df, os.path.join(plots_dir, plot_name), config)
        gc.collect()
    elapsed = time.time() - start_time
    if progress_callback is not None:
        try:
            progress_callback(len(traces), len(traces), "all traces", "writing outputs")
        except Exception:
            pass

    results_df = pd.DataFrame(results)
    if not results_df.empty:
        # Add an interpretation layer separate from metric-computation PASS/FAIL.
        quality_flags = [metric_quality_flag_for_result(r) for r in results_df.to_dict(orient="records")]
        review_flags = [manual_review_required_for_result(r) for r in results_df.to_dict(orient="records")]
        results_df.insert(0, "method_profile", config.method_profile)
        results_df.insert(1, "baseline_integrity", "PASS" if baseline_check["is_manuscript_baseline"] else "NON_MANUSCRIPT_MODIFIED")
        results_df.insert(2, "metric_quality_flag", quality_flags)
        results_df.insert(3, "manual_review_required", review_flags)
    group_df = summarize_groups(results_df) if config.export_group_summary else pd.DataFrame()
    validation_df = pd.DataFrame()
    if config.run_scotch_validation_if_detected and any("scotch" in str(s).lower() for s in results_df.get("sheet", [])):
        validation_df = validate_scotch(results_df, config.validation_relative_tolerance, config.validation_absolute_tolerance)

    # Readiness after run.
    pass_count = int((results_df.get("status", pd.Series(dtype=str)) == "PASS").sum()) if not results_df.empty else 0
    fail_count = int((results_df.get("status", pd.Series(dtype=str)) == "FAIL").sum()) if not results_df.empty else 0
    review_count = int(results_df.get("manual_review_required", pd.Series(dtype=bool)).astype(bool).sum()) if not results_df.empty and "manual_review_required" in results_df else 0
    readiness_after = pd.DataFrame([{
        "input_file": os.path.basename(config.input_path),
        "run_folder": outdir,
        "traces_attempted": len(results_df),
        "metric_pass_count": pass_count,
        "metric_fail_count": fail_count,
        "manual_review_required_count": review_count,
        "analysis_status": "FAIL" if fail_count else ("PASS_WITH_CAUTION" if review_count else "PASS"),
        "analysis_mode": "manuscript_baseline_v1 strict Top5+Bottom5",
        "method_profile": config.method_profile,
        "baseline_integrity": "PASS" if baseline_check["is_manuscript_baseline"] else "NON_MANUSCRIPT_MODIFIED",
        "method_profile_hash": baseline_check["method_profile_hash"],
        "strict_metric_rule": f">={config.min_window_points} window points and >={config.required_top_peaks} peaks + >={config.required_bottom_troughs} troughs",
        "elapsed_seconds": round(elapsed, 3),
        "estimated_seconds_per_trace": round(elapsed / max(len(results_df), 1), 3),
        "recommendation_note": "If traces failed, optional recovery/recommendation scan can be run separately; it is not part of manuscript replication.",
    }])

    # Exports.
    paper_metrics_file = os.path.join(outdir, "paper_metrics.csv")
    group_file = os.path.join(outdir, "group_summary.csv")
    readiness_after_file = os.path.join(outdir, "readiness_after_run.csv")
    formula_file = os.path.join(outdir, "formula_dictionary.csv")
    guide_file = os.path.join(outdir, "metric_guide.csv")
    output_guide_file = os.path.join(outdir, "output_file_guide.csv")
    input_audit_file = os.path.join(outdir, "input_audit.csv")
    validation_file = os.path.join(outdir, "scotch_validation.csv")
    output_consistency_file = os.path.join(outdir, "output_consistency_audit.csv")
    diagnostic_summary_file = os.path.join(outdir, "diagnostic_summary.csv")
    issue_action_file = os.path.join(outdir, "issue_action_table.csv")
    validation_summary_file = os.path.join(outdir, "validation_count_summary.csv")
    trace_integrity_file = os.path.join(outdir, "trace_integrity_audit.csv")
    run_history_file = os.path.join(outdir, "run_history.csv")
    provenance_file = os.path.join(outdir, "provenance.json")
    method_file = os.path.join(outdir, "method_profile_manuscript_baseline_v1_4_0_rc15.json")

    results_df.to_csv(paper_metrics_file, index=False)
    group_df.to_csv(group_file, index=False)
    readiness_after.to_csv(readiness_after_file, index=False)
    formula_dictionary().to_csv(formula_file, index=False)
    metric_guide().to_csv(guide_file, index=False)
    output_file_guide().to_csv(output_guide_file, index=False)
    mapping.to_csv(input_audit_file, index=False)
    trace_integrity_df.to_csv(trace_integrity_file, index=False)
    run_history_df = pd.DataFrame([{
        "run_id": 1,
        "run_type": config.method_profile,
        "profile_name": getattr(config, "profile_name", config.method_profile),
        "profile_hash": baseline_check["method_profile_hash"],
        "input_file": os.path.basename(config.input_path),
        "input_file_hash": sha256_file(config.input_path) if os.path.exists(config.input_path) else "",
        "mapping_file": mapping_file,
        "backend_version": __version__,
        "start_time": now_stamp(),
        "elapsed_seconds": round(elapsed, 3),
        "n_traces": int(len(results_df)),
        "n_pass": pass_count,
        "n_fail": fail_count,
        "n_review": review_count,
        "output_folder": outdir,
    }])
    run_history_df.to_csv(run_history_file, index=False)
    if not validation_df.empty:
        validation_df.to_csv(validation_file, index=False)

    method_profile = asdict(config)
    method_profile.update({
        "version": "v1.4.0-rc15 manuscript-baseline",
        "method_profile": config.method_profile,
        "baseline_integrity": "PASS" if baseline_check["is_manuscript_baseline"] else "NON_MANUSCRIPT_MODIFIED",
        "method_profile_hash": baseline_check["method_profile_hash"],
        "changed_manuscript_parameters": baseline_check["changed_parameters"],
        "locked_manuscript_profile": LOCKED_MANUSCRIPT_PROFILE,
        "drift_rule": "abs(m) * L_win / Fbar_win <= 0.25",
        "energy_outputs": "excluded from this notebook",
        "strict_extrema_rule": "Top5/Bottom5 metrics withheld unless >=5 peaks and >=5 troughs are detected",
    })
    with open(method_file, "w", encoding="utf-8") as f:
        json.dump(method_profile, f, indent=2)

    warning_records = []
    if not results_df.empty:
        for _, _r in results_df.iterrows():
            _warn = str(_r.get("warnings", ""))
            _fail = str(_r.get("failure_reason", ""))
            if _warn or _fail:
                warning_records.append({"sheet": _r.get("sheet", ""), "status": _r.get("status", ""), "warnings": _warn, "failure_reason": _fail})

    provenance = {
        "version": "v1.4.0-rc15 manuscript-baseline",
        "software_versions": software_versions(),
        "timestamp": now_stamp(),
        "input_path": config.input_path,
        "input_sha256": sha256_file(config.input_path) if os.path.exists(config.input_path) else "",
        "formulas": formula_dictionary().to_dict(orient="records"),
        "method_profile": config.method_profile,
        "baseline_integrity": "PASS" if baseline_check["is_manuscript_baseline"] else "NON_MANUSCRIPT_MODIFIED",
        "method_profile_hash": baseline_check["method_profile_hash"],
        "changed_manuscript_parameters": baseline_check["changed_parameters"],
        "warnings": warning_records,
        "outputs": {
            "paper_metrics": paper_metrics_file,
            "group_summary": group_file,
            "readiness_after_run": readiness_after_file,
            "formula_dictionary": formula_file,
            "metric_guide": guide_file,
            "output_file_guide": output_guide_file,
            "input_audit": input_audit_file,
            "scotch_validation": validation_file if not validation_df.empty else "",
            "output_consistency_audit": output_consistency_file,
            "diagnostic_summary": diagnostic_summary_file,
            "issue_action_table": issue_action_file,
            "validation_count_summary": validation_summary_file,
            "trace_integrity_audit": trace_integrity_file,
            "run_history": run_history_file,
            "run_report_html": report_html_file if 'report_html_file' in locals() else "",
            "run_report_markdown": report_md_file if 'report_md_file' in locals() else "",
            "qc_plots": plots_dir,
            "method_profile": method_file,
        },
        "config": method_profile,
    }
    with open(provenance_file, "w", encoding="utf-8") as f:
        json.dump(provenance, f, indent=2)

    excel_file = ""
    if config.export_excel:
        excel_file = os.path.join(outdir, "analysis_report.xlsx")
        with pd.ExcelWriter(excel_file, engine="openpyxl") as writer:
            readiness_after.to_excel(writer, sheet_name="Readiness", index=False)
            results_df.to_excel(writer, sheet_name="Paper_Metrics", index=False)
            group_df.to_excel(writer, sheet_name="Group_Summary", index=False)
            formula_dictionary().to_excel(writer, sheet_name="Formula_Dictionary", index=False)
            metric_guide().to_excel(writer, sheet_name="Metric_Guide", index=False)
            output_file_guide().to_excel(writer, sheet_name="Output_File_Guide", index=False)
            if not validation_df.empty:
                validation_df.to_excel(writer, sheet_name="Scotch_Validation", index=False)
            mapping.to_excel(writer, sheet_name="Input_Mapping", index=False)
            trace_integrity_df.to_excel(writer, sheet_name="Trace_Integrity", index=False)

    # Verify that exported CSV and Excel tables are consistent with the in-memory results.
    consistency_df = run_output_consistency_audit(results_df, group_df, paper_metrics_file, group_file, excel_file)
    consistency_df.to_csv(output_consistency_file, index=False)
    diagnostic_df = diagnostic_summary(results_df, readiness_after, consistency_df, validation_df)
    issue_action_df = issue_action_table(results_df)
    validation_summary_df = validation_count_summary(validation_df, consistency_df)
    diagnostic_df.to_csv(diagnostic_summary_file, index=False)
    issue_action_df.to_csv(issue_action_file, index=False)
    validation_summary_df.to_csv(validation_summary_file, index=False)
    report_paths = write_run_report(
        outdir=outdir,
        results_df=results_df,
        readiness_after=readiness_after,
        diagnostic_df=diagnostic_df,
        issue_action_df=issue_action_df,
        validation_summary_df=validation_summary_df,
        trace_integrity_df=trace_integrity_df,
        run_history_df=run_history_df,
        provenance=provenance,
    )
    report_html_file = report_paths.get("html", "")
    report_md_file = report_paths.get("markdown", "")
    if excel_file:
        try:
            with pd.ExcelWriter(excel_file, engine="openpyxl", mode="a", if_sheet_exists="replace") as writer:
                consistency_df.to_excel(writer, sheet_name="Output_Consistency", index=False)
                diagnostic_df.to_excel(writer, sheet_name="Diagnostic_Summary", index=False)
                issue_action_df.to_excel(writer, sheet_name="Issue_Action_Table", index=False)
                validation_summary_df.to_excel(writer, sheet_name="Validation_Summary", index=False)
                trace_integrity_df.to_excel(writer, sheet_name="Trace_Integrity", index=False)
        except Exception:
            pass
    consistency_fail_count = int((consistency_df.get("status", pd.Series(dtype=str)) == "FAIL").sum()) if not consistency_df.empty else 0
    provenance["outputs"]["output_consistency_audit"] = output_consistency_file
    provenance["output_consistency_status"] = "PASS" if consistency_fail_count == 0 else "FAIL"
    provenance["output_consistency_fail_count"] = consistency_fail_count
    with open(provenance_file, "w", encoding="utf-8") as f:
        json.dump(provenance, f, indent=2)

    zip_file = ""
    if config.make_zip:
        zip_file = os.path.join(outdir, "manuscript_baseline_outputs.zip")
        with zipfile.ZipFile(zip_file, "w", compression=zipfile.ZIP_DEFLATED) as z:
            for fp in [paper_metrics_file, group_file, readiness_after_file, formula_file, guide_file, output_guide_file, input_audit_file, trace_integrity_file, output_consistency_file, diagnostic_summary_file, issue_action_file, validation_summary_file, run_history_file, report_html_file if 'report_html_file' in locals() else '', report_md_file if 'report_md_file' in locals() else '', provenance_file, method_file, excel_file]:
                if fp and os.path.exists(fp):
                    z.write(fp, arcname=os.path.basename(fp))
            if not validation_df.empty and os.path.exists(validation_file):
                z.write(validation_file, arcname=os.path.basename(validation_file))
            if config.export_qc_plots and os.path.isdir(plots_dir):
                for name in os.listdir(plots_dir):
                    z.write(os.path.join(plots_dir, name), arcname=f"qc_plots/{name}")

    return {
        "outdir": outdir,
        "paper_metrics_file": paper_metrics_file,
        "group_summary_file": group_file,
        "readiness_after_run_file": readiness_after_file,
        "formula_dictionary_file": formula_file,
        "metric_guide_file": guide_file,
        "output_file_guide_file": output_guide_file,
        "input_audit_file": input_audit_file,
        "output_consistency_audit_file": output_consistency_file,
        "diagnostic_summary_file": diagnostic_summary_file,
        "issue_action_table_file": issue_action_file,
        "validation_count_summary_file": validation_summary_file,
        "trace_integrity_audit_file": trace_integrity_file,
        "run_history_file": run_history_file,
        "run_report_html_file": report_html_file if 'report_html_file' in locals() else "",
        "run_report_markdown_file": report_md_file if 'report_md_file' in locals() else "",
        "validation_file": validation_file if not validation_df.empty else "",
        "excel_report_file": excel_file,
        "zip_file": zip_file,
        "results": results_df,
        "group_summary": group_df,
        "validation": validation_df,
        "readiness_after_run": readiness_after,
        "sheet_mapping_file": mapping_file,
        "column_assignment_audit_file": column_audit_file,
        "advanced_column_pair_scores_file": scan.get("advanced_column_pair_scores_file", "") if "scan" in locals() else "",
        "output_consistency": consistency_df,
        "diagnostic_summary": diagnostic_df,
        "issue_action_table": issue_action_df,
        "validation_count_summary": validation_summary_df,
        "trace_integrity_audit": trace_integrity_df,
        "run_history": run_history_df,
        "readiness_report_file": readiness_file,
    }


# =============================================================================
# Diagnostic summaries and recommendation tables
# =============================================================================

def validation_count_summary(validation_df: pd.DataFrame, consistency_df: pd.DataFrame) -> pd.DataFrame:
    """Summarize reference-benchmark and export-consistency checks.

    The output-consistency count is intentionally dynamic: it depends on the number
    of audited result columns and exported tables in the current software version.
    The target is zero failed checks, not a fixed total such as 194 or 200.
    """
    rows: List[Dict[str, Any]] = []
    if validation_df is not None and not validation_df.empty:
        total = int(len(validation_df))
        passed = int((validation_df.get("validation_status", pd.Series(dtype=str)) == "PASS").sum())
        rows.append({
            "validation_block": "Reference benchmark validation",
            "checks_total": total,
            "checks_passed": passed,
            "checks_failed": total - passed,
            "meaning": "Bundled Scotch Tape reference-benchmark data reproduce expected row-level metric values within tolerance.",
            "why_it_matters": "Confirms that the analysis engine, formulas, reference data, and environment reproduce the package benchmark before user-data interpretation.",
            "what_it_does_not_mean": "It does not validate the user's specimen preparation, experimental design, or trace quality.",
            "count_note": "This block is fixed for the bundled reference file unless the benchmark file or expected-value file changes.",
        })
    if consistency_df is not None and not consistency_df.empty:
        total = int(len(consistency_df))
        passed = int((consistency_df.get("status", pd.Series(dtype=str)) == "PASS").sum())
        rows.append({
            "validation_block": "Output consistency audit",
            "checks_total": total,
            "checks_passed": passed,
            "checks_failed": total - passed,
            "meaning": "Exported CSV/Excel tables match the in-memory calculated result tables.",
            "why_it_matters": "Prevents silent export mismatches, dropped columns, row-count changes, and numeric changes during file writing.",
            "what_it_does_not_mean": "It does not mean every trace is scientifically clean; review metric-quality flags and QC plots.",
            "count_note": "The total count is dynamic because it includes row-count, column-set, and per-column value checks for each exported table. If the exported schema changes, this number changes. Zero failed checks is the critical condition.",
        })
    return pd.DataFrame(rows)


def issue_action_table(results: pd.DataFrame) -> pd.DataFrame:
    """Map common analysis warnings/failures to conservative next actions."""
    if results is None or results.empty:
        return pd.DataFrame([{
            "issue_detected": "No results available",
            "meaning": "No analysis table was produced.",
            "suggested_action": "Run the manuscript baseline first.",
            "caution": "Do not create modified settings before confirming input file, geometry, and columns.",
        }])
    fail_reason = results.get("failure_reason", pd.Series("", index=results.index)).fillna("").astype(str)
    warnings_col = results.get("warnings", pd.Series("", index=results.index)).fillna("").astype(str)
    trust = results.get("extrema_trust_flag", pd.Series("", index=results.index)).fillna("").astype(str)
    rows: List[Dict[str, Any]] = []
    def add(issue, meaning, action, caution, n):
        if n:
            rows.append({
                "issue_detected": issue,
                "affected_traces": int(n),
                "meaning": meaning,
                "suggested_action": action,
                "caution": caution,
            })
    add(
        "No valid analysis window",
        "A 25 mm manuscript window was not found after start/tail exclusion and drift filtering.",
        "Review trace length and displacement units. A shorter global window can be tested only in NON_MANUSCRIPT_MODIFIED recovery mode.",
        "Do not shorten the window merely to force a pass; shorter windows use less force-history information and can increase sensitivity to local noise.",
        fail_reason.str.contains("No valid", case=False, regex=False).sum(),
    )
    add(
        "Incomplete Top5/Bot5 extrema",
        "The selected window did not contain at least five peaks and five troughs.",
        "Inspect the QC plot. If the trace is truly smooth, do not force Top5/Bot5 metrics. If real oscillations are missed, a conservative prominence sensitivity check may be run.",
        "Lowering peak prominence can convert noise into artificial extrema.",
        fail_reason.str.contains("Incomplete extrema", case=False, regex=False).sum(),
    )
    add(
        "Clustered extrema",
        "Top5 peaks or Bottom5 troughs are concentrated in a narrow part of the selected window.",
        "Inspect QC plot manually. Optional strict extrema-spread QC can withhold these metrics in a recovery/sensitivity run.",
        "Manuscript baseline reports this as an advisory quality flag, not as an automatic failure.",
        trust.eq("LOW_EXTREMA_CLUSTERED").sum(),
    )
    add(
        "Very short or sparse trace",
        "There are too few processed numeric data points for the strict metric gate.",
        "Check instrument export, sampling interval, and column assignment before changing analysis settings.",
        "A modified analysis cannot recover missing data points.",
        fail_reason.str.contains("Too few", case=False, regex=False).sum(),
    )
    add(
        "Warnings present",
        "At least one trace contains non-fatal warnings.",
        "Read paper_metrics.csv and the QC plots before using the metric values in ranking or reporting.",
        "Warnings may indicate trace quality, column/unit, extrema, or break-proxy issues.",
        warnings_col.str.len().gt(0).sum(),
    )
    if not rows:
        rows.append({
            "issue_detected": "No major metric-gate issue detected",
            "affected_traces": 0,
            "meaning": "All analyzed traces passed the manuscript metric gate without major warnings.",
            "suggested_action": "Use manuscript-baseline outputs for paper-replica interpretation after reviewing QC plots.",
            "caution": "Export consistency confirms file-writing integrity; it does not replace visual QC review.",
        })
    return pd.DataFrame(rows)


def diagnostic_summary(results: pd.DataFrame, readiness_after: pd.DataFrame, consistency_df: pd.DataFrame, validation_df: pd.DataFrame) -> pd.DataFrame:
    """Create a compact pass/caution/fail summary for notebook display and export."""
    if readiness_after is not None and not readiness_after.empty:
        r = readiness_after.iloc[0].to_dict()
    else:
        r = {}
    consistency_fail = int((consistency_df.get("status", pd.Series(dtype=str)) == "FAIL").sum()) if isinstance(consistency_df, pd.DataFrame) and not consistency_df.empty else 0
    rows = [
        {"status_area": "Analysis status", "status": r.get("analysis_status", "UNKNOWN"), "meaning": "Metric gate result across analyzed user traces.", "next_action": "Review fail/caution rows and QC plots before interpreting rankings."},
        {"status_area": "Metric quality", "status": "REVIEW_REQUIRED" if int(r.get("manual_review_required_count", 0) or 0) else "QUALITY_OK", "meaning": "Whether computed metrics have advisory QC flags.", "next_action": "If review is required, inspect QC plots and issue_action_table.csv."},
        {"status_area": "Export integrity", "status": "FAIL" if consistency_fail else "PASS", "meaning": "Whether CSV/Excel exports match in-memory result tables.", "next_action": "If FAIL, do not use exported results until rerun or debugged."},
    ]
    return pd.DataFrame(rows)


def run_record_from_outputs(outputs: Dict[str, Any], config: ManuscriptConfig, run_label: str = "") -> Dict[str, Any]:
    """Return a compact record for run-history/run-comparison tables."""
    r = outputs.get("readiness_after_run") if isinstance(outputs, dict) else pd.DataFrame()
    row = r.iloc[0].to_dict() if isinstance(r, pd.DataFrame) and not r.empty else {}
    return {
        "run_label": run_label or row.get("method_profile", "run"),
        "method_profile": row.get("method_profile", getattr(config, "method_profile", "")),
        "baseline_integrity": row.get("baseline_integrity", ""),
        "window_length_mm": getattr(config, "window_length_mm", np.nan),
        "start_offset_mm": getattr(config, "start_offset_mm", np.nan),
        "tail_exclude_frac": getattr(config, "tail_exclude_frac", np.nan),
        "max_window_drift_frac": getattr(config, "max_window_drift_frac", np.nan),
        "peak_prominence_frac": getattr(config, "peak_prominence_frac", np.nan),
        "strict_extrema_spread_check": getattr(config, "strict_extrema_spread_check", False),
        "traces_attempted": row.get("traces_attempted", np.nan),
        "metric_pass_count": row.get("metric_pass_count", np.nan),
        "metric_fail_count": row.get("metric_fail_count", np.nan),
        "manual_review_required_count": row.get("manual_review_required_count", np.nan),
        "analysis_status": row.get("analysis_status", ""),
        "outdir": outputs.get("outdir", "") if isinstance(outputs, dict) else "",
        "zip_file": outputs.get("zip_file", "") if isinstance(outputs, dict) else "",
    }

# =============================================================================
# Optional recovery recommendation helpers
# =============================================================================

def suggest_global_settings_from_results(results: pd.DataFrame, config: ManuscriptConfig) -> pd.DataFrame:
    """Suggest a single global non-manuscript setting set from baseline diagnostics.

    The suggestion is conservative: it does not optimize adhesive rankings. It only
    responds to feasibility failures/warnings such as too-short traces, no valid
    25 mm window, excessive drift, or clustered extrema.
    """
    baseline = {
        "window_length_mm": float(LOCKED_MANUSCRIPT_PROFILE["window_length_mm"]),
        "start_offset_mm": float(LOCKED_MANUSCRIPT_PROFILE["start_offset_mm"]),
        "tail_exclude_frac": float(LOCKED_MANUSCRIPT_PROFILE["tail_exclude_frac"]),
        "max_window_drift_frac": float(LOCKED_MANUSCRIPT_PROFILE["max_window_drift_frac"]),
        "peak_prominence_frac": float(LOCKED_MANUSCRIPT_PROFILE["peak_prominence_frac"]),
        "strict_extrema_spread_check": False,
    }
    suggested = dict(baseline)
    reasons = {k: "unchanged; manuscript baseline remains feasible for this setting" for k in baseline}

    if results is None or results.empty:
        return pd.DataFrame([
            {"parameter": k, "manuscript_baseline": v, "suggested_global": v, "why": "No baseline results available."}
            for k, v in baseline.items()
        ])

    fail_reason = results.get("failure_reason", pd.Series("", index=results.index)).fillna("").astype(str)
    warnings_col = results.get("warnings", pd.Series("", index=results.index)).fillna("").astype(str)
    status = results.get("status", pd.Series("", index=results.index)).fillna("").astype(str)

    spans = pd.to_numeric(results.get("distance_max_mm", pd.Series(dtype=float)), errors="coerce")
    finite_spans = spans[np.isfinite(spans)]
    no_window = fail_reason.str.contains("No valid 25", case=False, regex=False).any()
    too_few = fail_reason.str.contains("Too few", case=False, regex=False).any()
    incomplete_extrema = fail_reason.str.contains("Incomplete extrema", case=False, regex=False).any()
    clustered = results.get("extrema_trust_flag", pd.Series("", index=results.index)).fillna("").astype(str).eq("LOW_EXTREMA_CLUSTERED").any()

    if not finite_spans.empty:
        # Conservative estimate of usable travel after start/tail guard. Use the lower quartile
        # so one long trace does not hide a short-trace problem.
        q25_span = float(np.nanquantile(finite_spans, 0.25))
        feasible_window = (1.0 - baseline["tail_exclude_frac"]) * q25_span - baseline["start_offset_mm"]
        if no_window or feasible_window < baseline["window_length_mm"]:
            floor_window = 0.5 * baseline["window_length_mm"]
            raw_proposed = math.floor(max(feasible_window * 0.85, floor_window))
            proposed = max(floor_window, min(baseline["window_length_mm"], raw_proposed))
            suggested["window_length_mm"] = float(proposed)
            reasons["window_length_mm"] = (
                f"Lower-quartile trace span suggests only about {feasible_window:.1f} mm usable travel "
                f"after start/tail guards. Suggested recovery uses ~85% of feasible lower-quartile travel, "
                f"but does not go below half of the manuscript 25 mm window ({floor_window:.1f} mm) unless the user manually enters advanced settings."
            )

    if no_window and suggested["window_length_mm"] == baseline["window_length_mm"]:
        suggested["window_length_mm"] = 15.0
        reasons["window_length_mm"] = "No valid 25 mm window occurred in at least one trace; 15 mm is suggested as a conservative recovery test."

    if incomplete_extrema:
        # Do not automatically lower prominence; it can convert noise into peaks. Keep unchanged and warn.
        reasons["peak_prominence_frac"] = (
            "Some traces lacked 5 peaks/5 troughs. Prominence is left unchanged to avoid converting noise into extrema; "
            "inspect QC plots before changing this manually."
        )

    if clustered:
        reasons["strict_extrema_spread_check"] = (
            "At least one PASS trace had clustered Top5/Bot5 extrema. Manuscript mode reports this as caution; "
            "strict mode can be enabled to withhold such metrics in exploratory runs."
        )

    if too_few:
        reasons["window_length_mm"] = reasons.get("window_length_mm", "") + " One or more traces have very few processed points; check sampling/export before relying on recovery settings."

    rows = []
    for k in baseline:
        rows.append({
            "parameter": k,
            "manuscript_baseline": baseline[k],
            "suggested_global": suggested[k],
            "changed": baseline[k] != suggested[k],
            "why": reasons[k],
        })
    return pd.DataFrame(rows)




def feasible_window_summary_from_mapping(mapping_df: pd.DataFrame, config: ManuscriptConfig) -> pd.DataFrame:
    """Summarize whether the active mapping can support the selected window length.

    Uses rough_displacement_span_mm from scan_workbook when available. This is a
    pre-analysis feasibility check: it estimates travel after the start-offset and
    tail-exclusion guards and warns before a user profile is activated.
    """
    if mapping_df is None or mapping_df.empty:
        return pd.DataFrame([{"metric": "feasible_window", "value": np.nan, "message": "No mapping table available."}])
    df = mapping_df.copy()
    if "include" in df.columns:
        try:
            df = df[df["include"].astype(bool)]
        except Exception:
            pass
    spans = pd.to_numeric(df.get("rough_displacement_span_mm", pd.Series(dtype=float)), errors="coerce")
    spans = spans[np.isfinite(spans)]
    baseline_window = float(LOCKED_MANUSCRIPT_PROFILE.get("window_length_mm", 25.0))
    start = float(getattr(config, "start_offset_mm", LOCKED_MANUSCRIPT_PROFILE.get("start_offset_mm", 5.0)))
    tail = float(getattr(config, "tail_exclude_frac", LOCKED_MANUSCRIPT_PROFILE.get("tail_exclude_frac", 0.20)))
    active_window = float(getattr(config, "window_length_mm", baseline_window))
    floor_window = 0.5 * baseline_window
    if spans.empty:
        return pd.DataFrame([{
            "metric": "feasible_window", "value": np.nan,
            "message": "rough_displacement_span_mm was not available; window feasibility will be checked during analysis.",
            "active_window_length_mm": active_window,
            "suggested_recovery_window_mm": np.nan,
            "half_manuscript_window_floor_mm": floor_window,
        }])
    feasible = (1.0 - tail) * spans - start
    feasible = feasible[np.isfinite(feasible)]
    if feasible.empty:
        min_f = med_f = max_f = np.nan
        q25 = np.nan
    else:
        min_f = float(np.nanmin(feasible))
        med_f = float(np.nanmedian(feasible))
        max_f = float(np.nanmax(feasible))
        q25 = float(np.nanquantile(feasible, 0.25))
    raw_suggested = 0.85 * q25 if np.isfinite(q25) else np.nan
    if np.isfinite(raw_suggested):
        suggested = max(floor_window, min(baseline_window, math.floor(raw_suggested)))
    else:
        suggested = np.nan
    if np.isfinite(min_f) and active_window > min_f:
        status = "CAUTION_WINDOW_EXCEEDS_MIN_FEASIBLE"
        message = f"Active window {active_window:.2f} mm exceeds the minimum estimated feasible window {min_f:.2f} mm after start/tail guards. Some sheets may fail window selection."
    elif np.isfinite(active_window) and active_window < floor_window:
        status = "ADVANCED_CAUTION_BELOW_HALF_MANUSCRIPT_WINDOW"
        message = f"Active window {active_window:.2f} mm is below half of the manuscript 25 mm window ({floor_window:.2f} mm). Treat as advanced/manual sensitivity analysis."
    else:
        status = "OK_OR_REVIEW"
        message = "Active window does not exceed the minimum estimated feasible window from the current mapping, or feasibility cannot be fully determined until analysis."
    return pd.DataFrame([
        {"metric": "min_estimated_feasible_window_mm", "value": min_f, "message": "Minimum estimated window that can fit after start/tail guards across included sheets."},
        {"metric": "median_estimated_feasible_window_mm", "value": med_f, "message": "Median estimated feasible window across included sheets."},
        {"metric": "max_estimated_feasible_window_mm", "value": max_f, "message": "Maximum estimated feasible window across included sheets."},
        {"metric": "active_window_length_mm", "value": active_window, "message": status},
        {"metric": "suggested_recovery_window_mm", "value": suggested, "message": "~85% of lower-quartile feasible travel, floored at half of the manuscript 25 mm window unless advanced/manual mode is used."},
        {"metric": "window_feasibility_message", "value": np.nan, "message": message},
    ])

# =============================================================================
# Optional recovery recommendation stub
# =============================================================================



def infeasible_sheets_from_mapping(mapping_df: pd.DataFrame, config: ManuscriptConfig) -> pd.DataFrame:
    """Return sheet-level estimated window feasibility after start/tail guards.

    This is a pre-analysis guide for user-profile selection. It uses the scanned
    rough_displacement_span_mm when available and does not replace the final
    metric gate in run_manuscript_pipeline.
    """
    if mapping_df is None or mapping_df.empty:
        return pd.DataFrame()
    df = mapping_df.copy()
    if "include" in df.columns:
        try:
            df = df[df["include"].astype(bool)].copy()
        except Exception:
            pass
    if "rough_displacement_span_mm" not in df.columns:
        return pd.DataFrame()
    baseline_window = float(LOCKED_MANUSCRIPT_PROFILE.get("window_length_mm", 25.0))
    start = float(getattr(config, "start_offset_mm", LOCKED_MANUSCRIPT_PROFILE.get("start_offset_mm", 5.0)))
    tail = float(getattr(config, "tail_exclude_frac", LOCKED_MANUSCRIPT_PROFILE.get("tail_exclude_frac", 0.20)))
    active = float(getattr(config, "window_length_mm", baseline_window))
    floor = 0.5 * baseline_window
    rows = []
    for _, row in df.iterrows():
        span = pd.to_numeric(pd.Series([row.get("rough_displacement_span_mm", np.nan)]), errors="coerce").iloc[0]
        feasible = (1.0 - tail) * float(span) - start if np.isfinite(span) else np.nan
        if not np.isfinite(feasible):
            status = "UNKNOWN"
            msg = "rough_displacement_span_mm unavailable; final feasibility checked during analysis."
        elif active > feasible:
            status = "WINDOW_TOO_LONG_FOR_SHEET"
            msg = f"active window {active:.2f} mm exceeds estimated feasible window {feasible:.2f} mm."
        elif active < floor:
            status = "ADVANCED_SHORT_WINDOW"
            msg = f"active window {active:.2f} mm is below half of manuscript window ({floor:.2f} mm)."
        else:
            status = "FEASIBLE_ESTIMATE"
            msg = "active window is within estimated feasible span for this sheet."
        rows.append({
            "sheet": row.get("sheet", ""),
            "include": row.get("include", True),
            "rough_displacement_span_mm": span,
            "estimated_feasible_window_mm": feasible,
            "active_window_length_mm": active,
            "status": status,
            "message": msg,
        })
    return pd.DataFrame(rows)

def failed_trace_recovery_advice(results: pd.DataFrame) -> pd.DataFrame:
    """Provide lightweight recovery advice without changing manuscript outputs."""
    if results.empty or "status" not in results:
        return pd.DataFrame()
    failed = results[results["status"] != "PASS"].copy()
    if failed.empty:
        return pd.DataFrame([{"message": "All traces passed manuscript-baseline metric gates. Recommendation scan is optional and not needed for manuscript replication."}])
    rows = []
    for _, r in failed.iterrows():
        reason = str(r.get("failure_reason", ""))
        advice = "Review raw trace and mapping."
        if "No valid 25.0 mm" in reason:
            advice = "The 25.0 mm manuscript window may not fit this trace. Optional recovery can test shorter windows, but those outputs should not replace manuscript-replica metrics unless disclosed."
        elif "Incomplete extrema" in reason:
            advice = "The selected window does not contain 5 peaks and 5 troughs. Inspect QC plot; a smoother trace may not support Top5/Bottom5 metrics under strict manuscript mode."
        elif "Too few" in reason:
            advice = "Trace has too few usable points. Check data export or instrument sampling."
        rows.append({"sheet": r.get("sheet", ""), "failure_reason": reason, "recovery_advice": advice})
    return pd.DataFrame(rows)


def print_run_summary(outputs: Dict[str, Any]) -> None:
    r = outputs.get("readiness_after_run")
    if isinstance(r, pd.DataFrame) and not r.empty:
        row = r.iloc[0]
        print("\nAnalysis complete")
        print(f"Run folder: {outputs.get('outdir')}")
        print(f"Traces attempted: {row.get('traces_attempted')}")
        print(f"PASS: {row.get('metric_pass_count')} | FAIL: {row.get('metric_fail_count')}")
    for key in ["paper_metrics_file", "group_summary_file", "excel_report_file", "zip_file"]:
        val = outputs.get(key)
        if val:
            print(f"{key}: {val}")


# =============================================================================
# User profiles and user benchmark helpers (v1.4.0-rc15)
# =============================================================================

USER_TUNABLE_PARAMETERS = [
    "start_offset_mm", "window_length_mm", "tail_exclude_frac", "max_window_drift_frac",
    "peak_prominence_frac", "min_peak_distance_points", "min_window_points",
    "required_top_peaks", "required_bottom_troughs", "drop_threshold_frac",
    "break_proxy_consecutive_points", "extrema_spread_warning_frac",
    "strict_extrema_spread_check",
]

USER_PROFILE_PARAMETER_GUIDE = [
    {"parameter": "window_length_mm", "manuscript_baseline": 25.0, "plain_role": "Length of the selected force window.", "when_to_change": "Short traces or no valid baseline window.", "caution": "Shorter windows use less force-history information and can increase sensitivity."},
    {"parameter": "start_offset_mm", "manuscript_baseline": 5.0, "plain_role": "Initial displacement excluded before window search.", "when_to_change": "Large toe/slack region before active peel.", "caution": "Too large an offset can discard real peel data."},
    {"parameter": "tail_exclude_frac", "manuscript_baseline": 0.20, "plain_role": "Final fraction of trace excluded from window search.", "when_to_change": "Premature terminal drops or end effects.", "caution": "Too high a value may leave too little usable trace."},
    {"parameter": "max_window_drift_frac", "manuscript_baseline": 0.25, "plain_role": "Rejects strongly ramping selected windows.", "when_to_change": "Exploratory sensitivity only.", "caution": "Loosening the gate can accept stretching/ramping instead of plateau-like peel."},
    {"parameter": "peak_prominence_frac", "manuscript_baseline": 0.003, "plain_role": "Peak/trough detection threshold as a fraction of global corrected force.", "when_to_change": "Too many noise peaks or too few detected real extrema.", "caution": "Changing prominence can change Fc, Fci, and SSA. Inspect QC plots."},
    {"parameter": "strict_extrema_spread_check", "manuscript_baseline": False, "plain_role": "Treats clustered Top5/Bot5 extrema as failure rather than caution.", "when_to_change": "Suspected false-positive extrema around one hump.", "caution": "This is not the manuscript baseline behavior."},
]


def user_profile_parameter_guide() -> pd.DataFrame:
    """Return the guide for user-tunable non-manuscript profile parameters."""
    return pd.DataFrame(USER_PROFILE_PARAMETER_GUIDE)


def create_user_method_profile_template(path: str = "user_method_profile_template.json") -> str:
    """Create an editable user-method-profile JSON template without changing the built-in baseline."""
    profile = dict(LOCKED_MANUSCRIPT_PROFILE)
    profile.update({
        "profile_id": "user_profile_example_non_manuscript",
        "profile_name": "user_profile_example_non_manuscript",
        "display_name": "Example non-manuscript user profile",
        "profile_note": "Explain why this profile exists and what dataset type it was created for.",
        "profile_type": "NON_MANUSCRIPT_MODIFIED",
        "software_version_created": __version__,
        "important_note": "Edit this copy for your dataset. Do not overwrite the built-in manuscript_baseline_v1 method profile.",
    })
    with open(path, "w", encoding="utf-8") as f:
        json.dump(profile, f, indent=2)
    return path


def load_user_method_profile(path: str) -> Dict[str, Any]:
    """Load a user profile JSON.

    Tunable parameters are returned together with lightweight metadata used by the
    notebook interface (profile_id, display_name, and profile_note). The metadata
    is intentionally ignored by the numerical engine, but preserved for provenance.
    """
    with open(path, "r", encoding="utf-8") as f:
        raw = json.load(f)
    profile: Dict[str, Any] = {}
    for k in USER_TUNABLE_PARAMETERS:
        if k in raw:
            profile[k] = raw[k]
    for k in ["profile_id", "profile_name", "display_name", "profile_note", "note", "reason_for_change", "profile_type", "software_version_created"]:
        if k in raw:
            profile[k] = raw[k]
    if "profile_id" not in profile and "profile_name" in profile:
        profile["profile_id"] = str(profile["profile_name"])
    if "profile_note" not in profile and "note" in profile:
        profile["profile_note"] = str(profile["note"])
    return profile


def save_method_profile_from_config(
    config: ManuscriptConfig,
    path: str = "user_method_profile.json",
    profile_name: str = "user_profile_non_manuscript",
    display_name: Optional[str] = None,
    note: Optional[str] = None,
) -> str:
    """Save tunable settings from a config as a reusable user profile.

    The built-in manuscript profile remains read-only. This helper always writes a
    separate JSON file for user-defined or suggested settings.
    """
    profile_id = str(profile_name).replace(" ", "_")
    profile = {k: getattr(config, k) for k in USER_TUNABLE_PARAMETERS if hasattr(config, k)}
    profile.update({
        "profile_id": profile_id,
        "profile_name": profile_id,
        "display_name": display_name or profile_id,
        "profile_note": note or "Use this profile as a reproducible starting point for future runs. It does not change the built-in manuscript baseline.",
        "profile_type": "MANUSCRIPT_BASELINE" if validate_against_locked_manuscript_profile(config)["is_manuscript_baseline"] else "NON_MANUSCRIPT_MODIFIED",
        "software_version_created": __version__,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "method_profile_hash_reference": method_profile_hash(),
        "important_note": "This profile is user-created and separate from the built-in manuscript_baseline_v1 profile.",
    })
    with open(path, "w", encoding="utf-8") as f:
        json.dump(profile, f, indent=2)
    return path


def apply_user_profile_to_config(config: ManuscriptConfig, profile: Dict[str, Any], allow_non_manuscript: bool = True) -> ManuscriptConfig:
    """Apply a loaded user profile to a ManuscriptConfig in memory."""
    for k, v in profile.items():
        if k in USER_TUNABLE_PARAMETERS and hasattr(config, k):
            setattr(config, k, v)
    config.allow_non_manuscript_run = bool(allow_non_manuscript)
    config.method_profile = "non_manuscript_modified"
    return config


def create_user_benchmark_from_results(
    results_df: pd.DataFrame,
    expected_path: str = "user_benchmark_expected.csv",
    metadata_path: str = "user_benchmark_metadata.json",
    selected_sheets: Optional[Sequence[str]] = None,
    metrics: Optional[Sequence[str]] = None,
    note: str = "User-created benchmark from a trusted run.",
) -> Dict[str, str]:
    """Save selected rows/metrics from a trusted run as a separate user benchmark.

    This never overwrites the bundled Scotch reference benchmark. It is intended for
    user/lab reproducibility checks in future sessions.
    """
    if results_df is None or results_df.empty:
        raise ValueError("No result table is available to save as a user benchmark.")
    df = results_df.copy()
    if selected_sheets:
        df = df[df["sheet"].astype(str).isin([str(s) for s in selected_sheets])].copy()
    if df.empty:
        raise ValueError("No selected benchmark rows remain after sheet filtering.")
    if metrics is None:
        metrics = [
            "fc_N", "fci_N", "fc_over_w_N_per_mm", "fci_over_w_N_per_mm", "window_cv_pct",
            "ssa_abs_N", "ssa_norm", "psi_mu_over_sigma", "window_start_mm", "window_end_mm",
            "peak_count", "trough_count", "Top_Points", "Bot_Points",
        ]
    keep = ["sheet"] + [m for m in metrics if m in df.columns]
    expected = df[keep].copy()
    expected.to_csv(expected_path, index=False)
    meta = {
        "benchmark_type": "USER_BENCHMARK_READ_ONLY_COPY",
        "software_version_created": __version__,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "expected_file": expected_path,
        "selected_sheet_count": int(len(expected)),
        "metrics": [m for m in metrics if m in df.columns],
        "note": note,
        "warning": "This user benchmark is separate from the bundled Scotch reference benchmark and should be version-controlled by the user if needed.",
    }
    with open(metadata_path, "w", encoding="utf-8") as f:
        json.dump(meta, f, indent=2)
    return {"expected_path": expected_path, "metadata_path": metadata_path}


def validate_user_benchmark(
    results_df: pd.DataFrame,
    expected_path: str = "user_benchmark_expected.csv",
    relative_tolerance: float = 0.015,
    absolute_tolerance: float = 0.015,
) -> pd.DataFrame:
    """Compare current result rows against a user-created benchmark CSV."""
    if results_df is None or results_df.empty:
        return pd.DataFrame([{"sheet": "", "metric": "", "validation_status": "FAIL", "details": "No current results available."}])
    if not os.path.exists(expected_path):
        return pd.DataFrame([{"sheet": "", "metric": "", "validation_status": "FAIL", "details": f"Benchmark file not found: {expected_path}"}])
    expected = pd.read_csv(expected_path)
    rows: List[Dict[str, Any]] = []
    if "sheet" not in expected.columns or "sheet" not in results_df.columns:
        return pd.DataFrame([{"sheet": "", "metric": "", "validation_status": "FAIL", "details": "Both current results and expected benchmark need a sheet column."}])
    current = results_df.set_index(results_df["sheet"].astype(str))
    for _, erow in expected.iterrows():
        sheet = str(erow.get("sheet", ""))
        if sheet not in current.index:
            rows.append({"sheet": sheet, "metric": "row", "validation_status": "FAIL", "details": "Sheet not found in current results."})
            continue
        crow = current.loc[sheet]
        if isinstance(crow, pd.DataFrame):
            crow = crow.iloc[0]
        for metric in [c for c in expected.columns if c != "sheet"]:
            ev = erow.get(metric, np.nan)
            cv = crow.get(metric, np.nan)
            e_num = pd.to_numeric(pd.Series([ev]), errors="coerce").iloc[0]
            c_num = pd.to_numeric(pd.Series([cv]), errors="coerce").iloc[0]
            if np.isfinite(e_num) or np.isfinite(c_num):
                diff = abs(float(c_num) - float(e_num)) if np.isfinite(e_num) and np.isfinite(c_num) else np.inf
                tol = absolute_tolerance + relative_tolerance * max(abs(float(e_num)) if np.isfinite(e_num) else 0.0, 1e-12)
                ok = bool(diff <= tol)
                details = f"expected={e_num:.6g}; current={c_num:.6g}; abs_diff={diff:.3g}; tolerance={tol:.3g}"
            else:
                ok = str(ev) == str(cv)
                details = "text values match" if ok else f"expected={ev}; current={cv}"
            rows.append({"sheet": sheet, "metric": metric, "validation_status": "PASS" if ok else "FAIL", "details": details})
    return pd.DataFrame(rows)


def run_comparison_summary_from_history(run_history: Sequence[Dict[str, Any]]) -> pd.DataFrame:
    """Return run history as a compact comparison table."""
    if not run_history:
        return pd.DataFrame()
    return pd.DataFrame(list(run_history))


# =============================================================================
# Compact run report writer (v1.4.0-rc15)
# =============================================================================

def _report_table(df: pd.DataFrame, max_rows: int = 40) -> str:
    if df is None or not isinstance(df, pd.DataFrame) or df.empty:
        return "<p><em>No rows available.</em></p>"
    show = df.head(max_rows).copy()
    return show.to_html(index=False, escape=False, border=0, classes="report-table")


def write_run_report(
    outdir: str,
    results_df: pd.DataFrame,
    readiness_after: pd.DataFrame,
    diagnostic_df: pd.DataFrame,
    issue_action_df: pd.DataFrame,
    validation_summary_df: pd.DataFrame,
    trace_integrity_df: pd.DataFrame,
    run_history_df: pd.DataFrame,
    provenance: Dict[str, Any],
) -> Dict[str, str]:
    """Write a compact HTML/Markdown report for a run.

    The report is meant for humans reviewing one run. It summarizes the active
    profile, validation status, trace diagnostics, and key output files without
    exposing every internal audit table in the notebook UI.
    """
    ensure_dir(outdir)
    html_path = os.path.join(outdir, "run_report.html")
    md_path = os.path.join(outdir, "run_report.md")
    profile = provenance.get("method_profile", "") if isinstance(provenance, dict) else ""
    integrity = provenance.get("baseline_integrity", "") if isinstance(provenance, dict) else ""
    timestamp = provenance.get("timestamp", now_stamp()) if isinstance(provenance, dict) else now_stamp()
    input_name = os.path.basename(provenance.get("input_path", "")) if isinstance(provenance, dict) else ""
    css = """
    <style>
    body{font-family:Arial,Helvetica,sans-serif;margin:24px;line-height:1.45;color:#263238;}
    h1,h2{color:#1177B2}.card{border-left:6px solid #1177B2;background:#f4f8ff;border-radius:10px;padding:12px 16px;margin:12px 0;}
    .warn{border-left-color:#DD6B00;background:#fff8e1}.fail{border-left-color:#b71c1c;background:#fff5f5}.ok{border-left-color:#2e7d32;background:#f4fbf5}
    .report-table{border-collapse:collapse;width:100%;font-size:12px;margin:8px 0 18px 0}.report-table th{background:#eef3f7;text-align:left}.report-table th,.report-table td{border:1px solid #d9e1e5;padding:6px;vertical-align:top;word-break:break-word;}
    code{background:#eef1f3;padding:1px 4px;border-radius:4px;}
    </style>
    """
    readiness_html = _report_table(readiness_after, 10)
    diag_html = _report_table(diagnostic_df, 20)
    issue_html = _report_table(issue_action_df, 30)
    validation_html = _report_table(validation_summary_df, 20)
    integrity_html = _report_table(trace_integrity_df, 30)
    run_history_html = _report_table(run_history_df, 20)
    metric_cols = [c for c in ["sheet", "status", "metric_quality_flag", "manual_review_required", "failure_reason", "fc_over_w_N_per_mm", "fci_over_w_N_per_mm", "psi_mu_over_sigma", "ssa_abs_N", "break_proxy_status"] if isinstance(results_df, pd.DataFrame) and c in results_df.columns]
    metrics_html = _report_table(results_df[metric_cols] if metric_cols else results_df, 50)
    html = f"""<!doctype html><html><head><meta charset='utf-8'><title>Peel trace run report</title>{css}</head><body>
    <h1>Peel Trace Evaluation Run Report</h1>
    <div class='card'><b>Input:</b> <code>{input_name}</code><br><b>Profile:</b> <code>{profile}</code><br><b>Baseline integrity:</b> <code>{integrity}</code><br><b>Timestamp:</b> <code>{timestamp}</code></div>
    <h2>Run readiness</h2>{readiness_html}
    <h2>Diagnostic summary</h2>{diag_html}
    <h2>Issue-action guidance</h2>{issue_html}
    <h2>Validation and export-audit summary</h2>{validation_html}
    <h2>Trace-integrity audit</h2>{integrity_html}
    <h2>Per-trace metric preview</h2>{metrics_html}
    <h2>Run history</h2>{run_history_html}
    <div class='card warn'><b>Interpretation note:</b> Manuscript-baseline outputs are protocol-defined descriptors under the specified test geometry and analysis settings. NON_MANUSCRIPT_MODIFIED runs are recovery/sensitivity outputs and should not be mixed with manuscript-baseline outputs without explicit disclosure.</div>
    </body></html>"""
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)
    md_lines = [
        "# Peel Trace Evaluation Run Report",
        f"- Input: `{input_name}`",
        f"- Profile: `{profile}`",
        f"- Baseline integrity: `{integrity}`",
        f"- Timestamp: `{timestamp}`",
        "",
        "Open `run_report.html` for wrapped tables and QC-review context.",
        "",
        "## Interpretation note",
        "Manuscript-baseline outputs are protocol-defined descriptors under the specified test geometry and analysis settings. NON_MANUSCRIPT_MODIFIED runs are recovery/sensitivity outputs and should not be mixed with manuscript-baseline outputs without explicit disclosure.",
    ]
    with open(md_path, "w", encoding="utf-8") as f:
        f.write("\n".join(md_lines))
    return {"html": html_path, "markdown": md_path}
